#!/usr/bin/env python3
"""Convert a Qualtrics QSF survey into a printable Word response form.

Usage: python qsf_to_paper.py survey.qsf [--config config.json] [--survey-config survey_config.json]
Requires: python-docx (pip install python-docx)

Paper pagination ignores QSF page breaks. Block names and new pages per block
are optional and off by default; questions stay together where possible.
"""

import argparse
import copy
import json
import math
import re
import sys
from html.parser import HTMLParser
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SINGLE_MARK = "◯"
MULTIPLE_MARK = "☐"


# null means use the adaptive layout's existing default for that setting.
DEFAULT_CONFIG = {
    "fonts": {"family": "Arial", "title_size": 17, "section_size": 12,
              "preserve_rich_text": True},
    "sections": {"show_block_names": False},
    "page": {"paper_size": "letter", "orientation": "portrait",
             "margin_top_inches": .7, "margin_bottom_inches": .7,
             "margin_left_inches": .8, "margin_right_inches": .8},
    "questions": {"font_size": 12, "bold": False, "numbering": "sequential",
                  "blank_lines_before": 2, "line_spacing": None,
                  "space_after_pt": 5, "page_break_before": False},
    "responses": {"font_size": 11, "minimum_font_size": 8, "symbol_size": 11,
                  "single_mark": SINGLE_MARK, "multiple_mark": MULTIPLE_MARK,
                  "layout": "auto", "multiple_columns_above": 10,
                  "max_columns": 2, "short_label_max_chars": 30,
                  "dropdown_write_in_at": 12, "row_height_inches": .25,
                  "field_row_height_inches": .34,
                  "form_preferred_row_height_inches": .45,
                  "large_list_banded_rows": True, "indent_inches": .5},
    "tables": {"line_spacing": 1, "cell_padding_inches": None,
               "banded_rows": False, "band_color": "#F7F7F7",
               "borders": {"enabled": False, "weight_pt": .5, "color": "#D9D9D9"},
               "row_bottom_border": {"enabled": False, "weight_pt": .5, "color": "#CCCCCC"}},
    "matrix": {"header_rotation": "auto", "rotate_at_columns": 5,
               "statement_font_size": None, "minimum_statement_font_size": None,
               "header_font_size": 10, "minimum_header_font_size": 9,
               "header_bold": True, "header_max_height_inches": None,
               "statement_column_width_inches": None,
               "preferred_row_height_inches": None, "minimum_row_height_inches": .2,
               "maximum_row_height_inches": .45, "line_spacing": 1.1,
               "cell_padding_inches": 14 / 1440,
               "banded_rows": True, "band_color": "#F7F7F7", "header_color": "#EAEAEA",
               "row_bottom_border": {"enabled": False, "weight_pt": .5, "color": "#CCCCCC"}},
    "write_in": {"inline_short_answers": True, "inline_prompt_max_chars": 110,
                 "small_underscore_count": 20, "large_lines": 3,
                 "large_line_spacing": 1.5, "text_lines": 3,
                 "line_height_inches": .3},
    "pagination": {"new_page_per_block": False, "keep_questions_together": True,
                   "new_page_for_large_matrices": True, "large_matrix_min_rows": 8,
                   "large_matrix_min_columns": 9, "large_matrix_min_height_inches": 4.5},
    "question_overrides": {}
}


def merge_config(base, patch):
    result = copy.deepcopy(base)
    for key, value in patch.items():
        result[key] = merge_config(result[key], value) if isinstance(value, dict) and isinstance(result.get(key), dict) else copy.deepcopy(value)
    return result


def validate_patch(patch, schema=DEFAULT_CONFIG, path=""):
    if not isinstance(patch, dict):
        raise ValueError(f"{path or 'Configuration'} must be a JSON object.")
    enums = {"paper_size": {"letter", "a4", "legal"}, "orientation": {"portrait", "landscape"},
             "numbering": {"off", "sequential", "survey"},
             "header_rotation": {"auto", "always", "never"},
             "layout": {"auto", "vertical", "horizontal"}}
    for key, value in patch.items():
        name = f"{path}.{key}" if path else key
        if key not in schema:
            raise ValueError(f"Unknown setting: {name}")
        if key == "question_overrides":
            if not isinstance(value, dict):
                raise ValueError("question_overrides must be an object keyed by QID.")
            allowed = {k: DEFAULT_CONFIG[k] for k in ("fonts", "questions", "responses", "tables", "matrix", "write_in", "pagination")}
            for qid, override in value.items():
                if not qid or not isinstance(qid, str):
                    raise ValueError("Question override keys must be nonempty question IDs.")
                if isinstance(override, dict) and isinstance(override.get("fonts"), dict) and set(override["fonts"]) - {"family", "preserve_rich_text"}:
                    raise ValueError(f"{name}.{qid}: title/section sizes are document-wide.")
                if isinstance(override, dict) and isinstance(override.get("pagination"), dict) and "new_page_per_block" in override["pagination"]:
                    raise ValueError(f"{name}.{qid}: new_page_per_block is document-wide.")
                validate_patch(override, allowed, f"{name}.{qid}")
        elif isinstance(schema[key], dict):
            validate_patch(value, schema[key], name)
        elif value is None:
            if schema[key] is not None:
                raise ValueError(f"{name} cannot be null.")
        elif isinstance(schema[key], bool):
            if type(value) is not bool:
                raise ValueError(f"{name} must be true or false.")
        elif key in enums:
            if not isinstance(value, str) or value not in enums[key]:
                raise ValueError(f"{name} must be one of {', '.join(sorted(enums[key]))}.")
        elif key.endswith("color") or key == "color":
            if not isinstance(value, str) or not re.fullmatch(r"#?[0-9a-fA-F]{6}", value):
                raise ValueError(f"{name} must be a six-digit hex color, e.g. #CCCCCC.")
        elif isinstance(schema[key], str):
            if not isinstance(value, str) or not value.strip():
                raise ValueError(f"{name} must be a nonempty string.")
        else:
            if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
                raise ValueError(f"{name} must be a finite number.")
            zero_allowed = key.startswith("margin_") or key in {"blank_lines_before", "space_after_pt", "cell_padding_inches", "multiple_columns_above", "indent_inches"}
            if value < 0 or (not zero_allowed and value == 0):
                raise ValueError(f"{name} must be {'nonnegative' if zero_allowed else 'positive'}.")
            integer_keys = {"blank_lines_before", "multiple_columns_above", "max_columns", "short_label_max_chars",
                            "dropdown_write_in_at", "rotate_at_columns", "inline_prompt_max_chars",
                            "small_underscore_count", "large_lines", "text_lines", "large_matrix_min_rows", "large_matrix_min_columns"}
            if key in integer_keys and not isinstance(value, int):
                raise ValueError(f"{name} must be an integer.")
            if "font_size" in key or key in {"symbol_size", "title_size", "section_size"}:
                if not 4 <= value <= 72:
                    raise ValueError(f"{name} must be between 4 and 72 pt.")
            if key == "weight_pt" and not .25 <= value <= 12:
                raise ValueError(f"{name} must be between 0.25 and 12 pt.")


def validate_resolved(config):
    r, m = config["responses"], config["matrix"]
    for low, high, label in [(r["minimum_font_size"], r["font_size"], "Response font sizes"),
                             (m["minimum_header_font_size"], m["header_font_size"], "Matrix header font sizes"),
                             (m["minimum_statement_font_size"] or r["minimum_font_size"], m["statement_font_size"] or r["font_size"], "Matrix statement font sizes"),
                             (m["minimum_row_height_inches"], m["maximum_row_height_inches"], "Matrix row heights")]:
        if low > high:
            raise ValueError(f"{label}: minimum cannot exceed preferred/maximum.")
    p = config["page"]
    width, height = {"letter": (8.5, 11), "a4": (210 / 25.4, 297 / 25.4), "legal": (8.5, 14)}[p["paper_size"]]
    if p["orientation"] == "landscape": width, height = height, width
    if width - p["margin_left_inches"] - p["margin_right_inches"] < 2 or height - p["margin_top_inches"] - p["margin_bottom_inches"] < 2:
        raise ValueError("Page margins must leave at least 2 inches of printable width and height.")
    if width - p["margin_left_inches"] - p["margin_right_inches"] - r["indent_inches"] < 1.2:
        raise ValueError("Response indentation must leave at least 1.2 inches for multiple-choice tables.")


def resolve_config(config=None, cli_overrides=None):
    patch = {} if config is None else config
    validate_patch(patch)
    validate_patch(cli_overrides or {})
    result = merge_config(merge_config(DEFAULT_CONFIG, patch), cli_overrides or {})
    validate_resolved(result)
    for override in result["question_overrides"].values():
        validate_resolved(merge_config(result, override))
    return result


def load_config_layers(main_path=None, survey_path=None):
    """Read optional JSON layers; missing files leave lower-priority settings intact.

    The automatic main config is beside this script, independent of the CWD.
    Explicit paths are relative to the caller's working directory.
    """
    merged, warnings = {}, []
    main = Path(main_path) if main_path is not None else Path(__file__).resolve().with_name("config.json")
    layers = [(main, main_path is not None)]
    if survey_path is not None: layers.append((Path(survey_path), True))
    for path, explicit in layers:
        try:
            text = path.read_text(encoding="utf-8-sig")
        except FileNotFoundError:
            if explicit:
                warnings.append(f"Config file {path} was not found; continuing with the available settings.")
            continue
        except OSError as exc:
            warnings.append(f"Cannot read config file {path}: {exc}; continuing with the available settings.")
            continue
        try:
            patch = json.loads(text)
            validate_patch(patch)
        except (ValueError, json.JSONDecodeError) as exc:
            raise ValueError(f"Config file {path}: {exc}") from exc
        merged = merge_config(merged, patch)
    return merged, warnings


def sizes_between(preferred, minimum):
    values = [preferred]
    candidate = math.ceil(preferred) - 1
    while candidate > minimum:
        values.append(candidate)
        candidate -= 1
    if minimum not in values: values.append(minimum)
    return values


def config_assignment(patch, assignment):
    if "=" not in assignment:
        raise ValueError("--set expects section.setting=value.")
    name, raw = assignment.split("=", 1)
    keys = name.split(".")
    if any(not key for key in keys): raise ValueError(f"Invalid setting path: {name}")
    try: value = json.loads(raw)
    except json.JSONDecodeError: value = raw
    target = patch
    for key in keys[:-1]:
        target = target.setdefault(key, {})
        if not isinstance(target, dict): raise ValueError(f"Conflicting --set path: {name}")
    target[keys[-1]] = value


class RichText(str):
    """A string for layout calculations plus styled spans for Word runs."""

    def __new__(cls, spans):
        obj = super().__new__(cls, ''.join(text for text, _ in spans))
        obj.spans = spans
        return obj

    def __add__(self, other):
        return RichText(self.spans + getattr(other, 'spans', [(str(other), {})]))


def normalized_rich(spans, flatten=False):
    chars = []
    for text, style in spans:
        for char in text:
            char = ' ' if char in '\t\r' or (flatten and char == '\n') else char
            if char == ' ' and (not chars or chars[-1][0] in {' ', '\n'}):
                continue
            if char == '\n':
                while chars and chars[-1][0] == ' ':
                    chars.pop()
                if not chars or chars[-1][0] == '\n':
                    continue
            chars.append((char, style))
    while chars and chars[-1][0].isspace():
        chars.pop()
    result = []
    for char, style in chars:
        if result and result[-1][1] == style:
            result[-1] = (result[-1][0] + char, style)
        else:
            result.append((char, style.copy()))
    return RichText(result)


def flat(value):
    return normalized_rich(getattr(value, 'spans', [(str(value), {})]), flatten=True)


def rich_runs(p, value, size=11, bold=False, font="Arial", preserve=True, force_bold=False):
    for text, style in getattr(value, 'spans', [(str(value), {})]):
        run = p.add_run(text)
        style = style if preserve else {}
        run.font.name = font
        run.font.size = Pt(size)
        run.bold = True if force_bold else style.get('bold', bold)
        run.italic = style.get('italic', False)
        run.underline = style.get('underline', False)
        run.font.strike = style.get('strike', False)
        if style.get('super'):
            run.font.superscript = True
        elif style.get('sub'):
            run.font.subscript = True
        color = style.get('color', '')
        if re.fullmatch(r'[0-9a-fA-F]{6}', color):
            run.font.color.rgb = RGBColor.from_string(color)


class HTMLText(HTMLParser):
    """Conservatively retain visible text, breaks, and image alt text."""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts = []
        self.spans = []
        self.stack = []
        self.style = {}
        self.hidden = 0

    def emit(self, text):
        if not self.hidden:
            self.parts.append(text)
            self.spans.append((text, self.style.copy()))

    def handle_starttag(self, tag, attrs):
        attributes = dict(attrs)
        if tag in {'script', 'style'}:
            self.hidden += 1
        if tag not in {'br', 'img', 'hr', 'input', 'meta', 'link'}:
            self.stack.append((tag, self.style.copy()))
            self.style = self.style.copy()
        for tags, key in [({'b', 'strong'}, 'bold'), ({'i', 'em'}, 'italic'),
                          ({'u'}, 'underline'), ({'s', 'strike', 'del'}, 'strike'),
                          ({'sup'}, 'super'), ({'sub'}, 'sub')]:
            if tag in tags:
                self.style[key] = True
        css = dict((key.strip().lower(), val.strip().lower())
                   for key, val in re.findall(r'([\w-]+)\s*:\s*([^;]+)', attributes.get('style', '')))
        weight = css.get('font-weight', '')
        if weight:
            self.style['bold'] = weight in {'bold', 'bolder'} or (weight.isdigit() and int(weight) >= 600)
        if 'font-style' in css:
            self.style['italic'] = css['font-style'] in {'italic', 'oblique'}
        if 'text-decoration' in css:
            self.style['underline'] = 'underline' in css['text-decoration']
            self.style['strike'] = 'line-through' in css['text-decoration']
        color = css.get('color', attributes.get('color', '')).lower()
        colors = {'black':'000000', 'red':'FF0000', 'blue':'0000FF', 'green':'008000', 'gray':'808080', 'white':'FFFFFF'}
        rgb = re.fullmatch(r'rgb\(\s*(\d+)\s*,\s*(\d+)\s*,\s*(\d+)\s*\)', color)
        if rgb:
            color = ''.join(f'{min(255, int(v)):02X}' for v in rgb.groups())
        elif color.startswith('#'):
            color = color[1:]
            if len(color) == 3:
                color = ''.join(c * 2 for c in color)
        else:
            color = colors.get(color, color)
        if re.fullmatch(r'[0-9a-fA-F]{6}', color):
            self.style['color'] = color
        if tag in {"br", "p", "div", "li", "tr", "h1", "h2", "h3"}:
            self.emit("\n")
        if tag == "li":
            self.emit("• ")
        if tag == "img":
            self.emit(attributes.get("alt") or "")

    def handle_endtag(self, tag):
        if tag in {"p", "div", "li", "tr", "h1", "h2", "h3"}:
            self.emit("\n")
        for i in range(len(self.stack)-1, -1, -1):
            if self.stack[i][0] == tag:
                self.style = self.stack[i][1]
                del self.stack[i:]
                break
        if tag in {'script', 'style'}:
            self.hidden = max(0, self.hidden - 1)

    def handle_data(self, data):
        self.emit(data)


def clean(value):
    if isinstance(value, dict):
        value = value.get("Display", value.get("Text", value.get("Label", value.get("Description", ""))))
    if value is None:
        return ""
    parser = HTMLText()
    parser.feed(str(value))
    return normalized_rich(parser.spans)


def ordered(mapping, order):
    if not isinstance(mapping, dict):
        return []
    keys = [str(x) for x in order] if isinstance(order, list) else []
    keys += [str(x) for x in mapping if str(x) not in keys]
    return [(key, mapping[key]) for key in keys if key in mapping]


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def label_size(values, width, bold=False, preferred=11, minimum=8):
    """Choose one size for a whole set of labels, never per option."""
    longest = max((len(word) for value in values for word in value.split()), default=0)
    available = max(.2, width - .16)
    for size in sizes_between(preferred, minimum):
        if longest * size * (.65 if bold else .56) / 72 <= available or size == minimum:
            return size


def paragraph_zero(p):
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    if p.paragraph_format.line_spacing is None:
        p.paragraph_format.line_spacing = 1


def text_width(text, size=11, bold=False):
    """Approximate Arial advances in points without another dependency."""
    lower_widths = dict(zip('abcdefghijklmnopqrstuvwxyz',
                           (.556,.556,.5,.556,.556,.278,.556,.556,.222,.222,.5,.222,.833,.556,.556,.556,.556,.333,.5,.278,.556,.5,.722,.5,.5,.5)))
    narrow = "ijlI.,:;!'|"
    wide = 'MWmw@'
    units = sum(lower_widths[c] if c in lower_widths else .28 if c in narrow else .85 if c in wide else
                .28 if c == ' ' else 1 if ord(c) > 255 else
                .67 if c.isupper() else .56 for c in text)
    return units * size * (1.06 if bold else 1)


def inline_blank(value, count=20):
    """Insert a blank after the first prompt line, before any instructions."""
    result, inserted = [], False
    for text, style in getattr(value, 'spans', [(str(value), {})]):
        if not inserted and '\n' in text:
            first, rest = text.split('\n', 1)
            result.extend([(first, style), (' ' + '_' * count, {}), ('\n' + rest, style)])
            inserted = True
        else:
            result.append((text, style))
    if not inserted:
        result.append((' ' + '_' * count, {}))
    return RichText(result)


def cell_text(cell, value, bold=False, size=11, font="Arial", preserve=True, line_spacing=1):
    cell.text = ""
    rich_runs(cell.paragraphs[0], value, size, bold, font, preserve)
    paragraph_zero(cell.paragraphs[0])
    cell.paragraphs[0].paragraph_format.line_spacing = line_spacing
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def estimated_lines(text, width, size=11):
    """Conservative Arial word-wrap estimate, including wide Unicode text."""
    capacity = max(1, max(.15, width - .16) * 72)
    total = 0
    for line in text.split('\n'):
        used, count = 0, 1
        for word in line.split():
            length = text_width(word, size, bold=True) * 1.12
            space = size * .28
            if used and used + space + length > capacity:
                count += 1
                used = 0
            if length > capacity:
                pieces = math.ceil(length / capacity)
                count += pieces - 1
                used = length - (pieces - 1) * capacity
            else:
                used += length + (space if used else 0)
        total += count
    return max(1, total)


def other_size(choice):
    """QSF exports use TextEntrySize or TextEntryLength for this setting."""
    if not isinstance(choice, dict) or str(choice.get('TextEntry', '')).lower() not in {'true', '1', 'on'}:
        return None
    value = str(choice.get('TextEntrySize') or choice.get('TextEntryLength') or 'Small').lower()
    return value if value in {'small', 'medium', 'large'} else 'small'


def no_split(row):
    trpr = row._tr.get_or_add_trPr()
    trpr.append(OxmlElement("w:cantSplit"))


def repeat_header(row):
    row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))


def bottom_line(cell, weight=.5, color="999999"):
    borders = cell._tc.get_or_add_tcPr().find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        cell._tc.get_or_add_tcPr().append(borders)
    for existing in list(borders.findall(qn("w:bottom"))):
        borders.remove(existing)
    border = OxmlElement("w:bottom")
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), str(round(weight * 8)))
    border.set(qn("w:color"), color.lstrip("#"))
    borders.append(border)


def is_rating_scale(q, labels):
    """Recognize anchored rating scales; nominal choices remain vertical."""
    if len(labels) < 2 or len(labels) > 11 or any(other for _, other in labels):
        return False
    values = [re.sub(r"[^a-z0-9 ]", "", label.lower()).strip() for label, _ in labels]
    # Qualtrics stores its predefined Yes/No list as a Likert Autoscale too.
    # Binary choices are ordinary vertical lists, regardless of that metadata.
    if {"yes", "no"}.issubset(values):
        return False
    config = q.get("Configuration", {})
    if isinstance(config, dict) and config.get("Autoscale"):
        autoscale = config["Autoscale"]
        yscale = autoscale.get("YScale", {}) if isinstance(autoscale, dict) else {}
        if isinstance(yscale, dict) and str(yscale.get("Name", "")).lower() == "yesno":
            return False
        return True
    for family in (("agree",), ("satisfied",), ("useful",), ("comfortable",),
                   ("likely",), ("important",), ("confident",), ("easy", "difficult"),
                   ("acceptable",), ("effective",), ("good", "bad"), ("excellent", "poor")):
        if all(any(stem in value for stem in family) or value in {"neutral", "neither", "not applicable", "na"}
               for value in values):
            return True
    prompt = clean(q.get("QuestionText", "")).lower()
    return (all(re.fullmatch(r"\d+", value) for value in values) and
            bool(re.search(r"\brate\b|\brating\b|\bon a scale\b", prompt)))


class Writer:
    def __init__(self, title, text_lines=None, include_ids=False, config=None):
        self.base_config = resolve_config(config)
        if text_lines is not None:
            self.base_config["write_in"]["text_lines"] = text_lines
        if include_ids:
            self.base_config["questions"]["numbering"] = "survey"
        validate_patch(self.base_config)
        validate_resolved(self.base_config)
        self.settings = self.base_config
        self.doc = Document()
        self.warnings = []
        self.number = 0
        self.flowing_tables = set()
        sec = self.doc.sections[0]
        page = self.settings["page"]
        width, height = {"letter": (8.5, 11), "a4": (210 / 25.4, 297 / 25.4), "legal": (8.5, 14)}[page["paper_size"]]
        if page["orientation"] == "landscape": width, height = height, width
        sec.orientation = WD_ORIENT.LANDSCAPE if page["orientation"] == "landscape" else WD_ORIENT.PORTRAIT
        sec.page_width, sec.page_height = Inches(width), Inches(height)
        sec.top_margin = Inches(page["margin_top_inches"])
        sec.bottom_margin = Inches(page["margin_bottom_inches"])
        sec.left_margin = Inches(page["margin_left_inches"])
        sec.right_margin = Inches(page["margin_right_inches"])
        self.width = (sec.page_width - sec.left_margin - sec.right_margin) / 914400
        normal = self.doc.styles["Normal"]
        normal.font.name = self.settings["fonts"]["family"]
        normal.font.size = Pt(self.settings["responses"]["font_size"])
        normal.font.color.rgb = RGBColor(0, 0, 0)
        normal.paragraph_format.space_after = Pt(5)
        for name, size in [("Title", self.settings["fonts"]["title_size"]), ("Heading 1", self.settings["fonts"]["section_size"])]:
            style = self.doc.styles[name]
            style.font.name = self.settings["fonts"]["family"]
            style.font.size = Pt(size)
            style.font.bold = True
            style.font.color.rgb = RGBColor(0, 0, 0)
            style.paragraph_format.space_before = Pt(12)
            style.paragraph_format.space_after = Pt(6)
        # Word's built-in Title style can inherit a colored bottom rule.
        title_ppr = self.doc.styles["Title"].element.get_or_add_pPr()
        borders = title_ppr.find(qn("w:pBdr"))
        if borders is not None:
            title_ppr.remove(borders)
        p = self.doc.add_paragraph(style="Title")
        p.add_run(title)

    @property
    def text_lines(self):
        return self.settings["write_in"]["text_lines"]

    def cell_text(self, cell, value, bold=False, size=None):
        cell_text(cell, value, bold, self.settings["responses"]["font_size"] if size is None else size,
                  self.settings["fonts"]["family"], self.settings["fonts"]["preserve_rich_text"],
                  self.settings["tables"]["line_spacing"])

    def label_size(self, values, width, bold=False):
        r = self.settings["responses"]
        return label_size(values, width, bold, r["font_size"], r["minimum_font_size"])

    def decorate_table(self, table, matrix=False, banded_rows=None):
        t = self.settings["tables"]
        cfg = self.settings["matrix"] if matrix else t
        padding = cfg["cell_padding_inches"]
        if padding is not None:
            existing = table._tbl.tblPr.find(qn("w:tblCellMar"))
            if existing is not None: table._tbl.tblPr.remove(existing)
            margins = OxmlElement("w:tblCellMar")
            for edge in ("top", "bottom", "left", "right"):
                node = OxmlElement("w:" + edge)
                node.set(qn("w:w"), str(round(padding * 1440)))
                node.set(qn("w:type"), "dxa")
                margins.append(node)
            table._tbl.tblPr.append(margins)
        existing = table._tbl.tblPr.find(qn("w:tblBorders"))
        if existing is not None: table._tbl.tblPr.remove(existing)
        borders = OxmlElement("w:tblBorders")
        b = t["borders"]
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            border = OxmlElement("w:" + edge)
            border.set(qn("w:val"), "single" if b["enabled"] else "nil")
            if b["enabled"]:
                border.set(qn("w:sz"), str(round(b["weight_pt"] * 8)))
                border.set(qn("w:color"), b["color"].lstrip("#"))
            borders.append(border)
        table._tbl.tblPr.append(borders)
        for n, row in enumerate(table.rows):
            if matrix and n == 0: continue
            band = n % 2 == 0 if matrix else n % 2 == 1
            for cell in row.cells:
                if (cfg["banded_rows"] if banded_rows is None else banded_rows) and band:
                    shade(cell, cfg["band_color"].lstrip("#"))
                border = cfg["row_bottom_border"]
                if border["enabled"]: bottom_line(cell, border["weight_pt"], border["color"])

    def keep_question(self, start):
        # A keep-with-next chain joins the prompt to every response-table row.
        # Word inserts the page break above the group when it fits on a page.
        for element in list(self.doc.element.body)[start:]:
            if element.tag == qn("w:p"):
                from docx.text.paragraph import Paragraph
                p = Paragraph(element, self.doc)
                p.paragraph_format.keep_together = self.settings["pagination"]["keep_questions_together"]
                p.paragraph_format.keep_with_next = self.settings["pagination"]["keep_questions_together"]
            elif element.tag == qn("w:tbl"):
                from docx.table import Table
                table = Table(element, self.doc)
                flowing = element in self.flowing_tables or not self.settings["pagination"]["keep_questions_together"]
                for i, row in enumerate(table.rows):
                    oversize = flowing and row.height is not None and row.height.pt > 600
                    if not oversize:
                        no_split(row)
                    for cell in row.cells:
                        if cell._tc.get_or_add_tcPr().find(qn('w:textDirection')) is None:
                            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                        for j, p in enumerate(cell.paragraphs):
                            paragraph_zero(p)
                            p.paragraph_format.keep_together = not oversize
                            p.paragraph_format.keep_with_next = (i == 0 and not oversize) if flowing else (i < len(table.rows)-1 or j < len(cell.paragraphs)-1)
        elements = [e for e in list(self.doc.element.body)[start:] if e.tag != qn('w:sectPr')]
        if elements and elements[-1].tag == qn('w:p'):
            from docx.text.paragraph import Paragraph
            Paragraph(elements[-1], self.doc).paragraph_format.keep_with_next = False

    def response_table(self, widths, rows, indent_inches=0, banded_rows=None):
        table = self.doc.add_table(rows=rows, cols=len(widths))
        table.autofit = False
        # Reserve the indentation in the column widths so the right edge still
        # stays within the printable page. Explicit widths also help LibreOffice.
        table_width = table._tbl.tblPr.find(qn("w:tblW"))
        table_width.set(qn("w:w"), str(round(sum(widths) * 1440)))
        table_width.set(qn("w:type"), "dxa")
        indent = OxmlElement("w:tblInd")
        indent.set(qn("w:w"), str(round(indent_inches * 1440)))
        indent.set(qn("w:type"), "dxa")
        table._tbl.tblPr.append(indent)
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            border = OxmlElement("w:" + edge)
            border.set(qn("w:val"), "nil")
            borders.append(border)
        table._tbl.tblPr.append(borders)
        for col, width in zip(table.columns, widths):
            col.width = Inches(width)
        for row in table.rows:
            no_split(row)
            row.height = Inches(self.settings["responses"]["row_height_inches"])
            row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
            for cell, width in zip(row.cells, widths):
                cell.width = Inches(width)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                paragraph_zero(cell.paragraphs[0])
                cell.paragraphs[0].paragraph_format.line_spacing = self.settings["tables"]["line_spacing"]
        self.decorate_table(table, banded_rows=banded_rows)
        return table

    def lines(self, count=None, indent_inches=0):
        table = self.response_table([self.width - indent_inches], count or self.text_lines,
                                    indent_inches=indent_inches)
        for row in table.rows:
            row.height = Inches(self.settings["write_in"]["line_height_inches"])
            if not self.settings["tables"]["row_bottom_border"]["enabled"]:
                bottom_line(row.cells[0])

    def remaining_page_height(self, end):
        """Estimate space before a question; Word remains the layout authority."""
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        sec = self.doc.sections[0]
        capacity = (sec.page_height - sec.top_margin - sec.bottom_margin) / 914400 * 72
        used, group = 0, 0

        def paragraph_height(p, width):
            inherited_size = p.style.font.size or self.doc.styles["Normal"].font.size
            size = max((r.font.size.pt for r in p.runs if r.font.size is not None),
                       default=inherited_size.pt if inherited_size is not None else self.base_config["responses"]["font_size"])
            spacing = p.paragraph_format.line_spacing
            line = spacing.pt if hasattr(spacing, "pt") else size * (spacing or 1.2)
            before = p.paragraph_format.space_before
            after = p.paragraph_format.space_after
            if before is None: before = p.style.paragraph_format.space_before
            if after is None: after = p.style.paragraph_format.space_after
            return estimated_lines(p.text, width, size) * line + (before.pt if before else 0) + (after.pt if after else 0)

        def flush():
            nonlocal used, group
            if used + group > capacity:
                used = group if group <= capacity else group % capacity
            else:
                used += group
            group = 0

        for element in list(self.doc.element.body)[:end]:
            if element.tag == qn("w:p"):
                p = Paragraph(element, self.doc)
                if p.paragraph_format.page_break_before:
                    flush()
                    used = 0
                group += paragraph_height(p, self.width)
                keep = p.paragraph_format.keep_with_next
                if keep is None: keep = p.style.paragraph_format.keep_with_next
            elif element.tag == qn("w:tbl"):
                table = Table(element, self.doc)
                for row in table.rows:
                    height = row.height.pt if row.height is not None else 0
                    for cell in row.cells:
                        if cell._tc.get_or_add_tcPr().find(qn("w:textDirection")) is not None:
                            continue
                        width = cell.width.inches if cell.width is not None else self.width / len(row.cells)
                        height = max(height, sum(paragraph_height(p, width) for p in cell.paragraphs))
                    group += height
                keep = any(p.paragraph_format.keep_with_next for cell in table.rows[-1].cells for p in cell.paragraphs)
            else:
                continue
            if not keep: flush()
        flush()
        return max(0, capacity - used)

    def fields(self, labels, answer_first=False, available_height=None, current_page_height=None):
        labels = [flat(label) for label in labels]
        widths = [.55, self.width - .55] if answer_first else [self.width * 3.4 / 6.9, self.width * 3.5 / 6.9]
        table = self.response_table(widths, len(labels))
        size = self.label_size(labels, widths[1] if answer_first else widths[0])
        heights = [self.settings["responses"]["field_row_height_inches"]] * len(labels)
        if available_height is not None and labels:
            # Form rows get extra handwriting space only within the height
            # available to a whole question on one page. Wrapped labels retain
            # their necessary height; row heights remain minimums, never caps.
            label_width = widths[1] if answer_first else widths[0]
            padding = self.settings["tables"]["cell_padding_inches"]
            padding = .04 if padding is None else padding
            heights = [max(h, estimated_lines(label, label_width, size) * size *
                           self.settings["tables"]["line_spacing"] / 72 + 2 * padding)
                       for h, label in zip(heights, labels)]
            preferred = self.settings["responses"]["form_preferred_row_height_inches"]
            growth = [max(0, preferred - h) for h in heights]
            if current_page_height is not None and sum(heights) * 72 <= current_page_height:
                available_height = min(available_height, current_page_height)
            extra = max(0, available_height / 72 - sum(heights))
            fraction = min(1, extra / sum(growth)) if sum(growth) else 0
            heights = [h + g * fraction for h, g in zip(heights, growth)]
        for row, label, height in zip(table.rows, labels, heights):
            answer, text = (row.cells[0], row.cells[1]) if answer_first else (row.cells[1], row.cells[0])
            self.cell_text(text, label, size=size)
            text.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            text.paragraphs[0].paragraph_format.space_after = Pt(0)
            if not self.settings["tables"]["row_bottom_border"]["enabled"]:
                bottom_line(answer)
            row.height = Inches(height)

    def options(self, choices, multiple=False, rating=False):
        choices = [(flat(label), other) for label, other in choices]
        cfg = self.settings["responses"]
        indent = cfg["indent_inches"]
        width = self.width - indent
        mark = cfg["multiple_mark"] if multiple else cfg["single_mark"]
        if cfg["layout"] != "auto": rating = cfg["layout"] == "horizontal" and not any(other for _, other in choices)
        if not choices:
            self.lines(1, indent_inches=indent)
            return
        banded = True if len(choices) > cfg["multiple_columns_above"] and cfg["large_list_banded_rows"] else None
        if rating:
            table = self.response_table([width / len(choices)] * len(choices), 2,
                                        indent_inches=indent, banded_rows=banded)
            size = self.label_size([label for label, _ in choices], width / len(choices))
            for i, (label, _) in enumerate(choices):
                self.cell_text(table.cell(0, i), label, size=size)
                self.cell_text(table.cell(1, i), mark, size=cfg["symbol_size"])
                for row in table.rows:
                    row.cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            for cell in table.rows[0].cells:
                cell.paragraphs[0].paragraph_format.keep_with_next = True
            return
        max_length = max((len(label) for label, _ in choices), default=0)
        count = len(choices)
        # Use two vertical columns only for long lists of short, plain labels.
        split = count > cfg["multiple_columns_above"] and max_length <= cfg["short_label_max_chars"] and not any(other or "\n" in label for label, other in choices)
        groups = min(cfg["max_columns"], max(1, math.ceil(count / max(1, cfg["multiple_columns_above"])))) if split else 1
        groups = min(groups, max(1, int(width / 1.2)))
        rows = math.ceil(count / groups)
        label_width = (width - (groups - 1) * .4) / groups - .3
        widths = []
        for group in range(groups):
            if group: widths.append(.4)
            widths.extend([.3, label_width])
        table = self.response_table(widths, rows, indent_inches=indent, banded_rows=banded)
        size = self.label_size([label for label, _ in choices], widths[1])
        for i, (label, other) in enumerate(choices):
            group, row = divmod(i, rows)
            mark_cell = table.cell(row, 3 * group)
            label_cell = table.cell(row, 3 * group + 1)
            self.cell_text(mark_cell, mark, size=cfg["symbol_size"])
            mark_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            mark_cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            # Write-in underscores remain in the same editable cell as the label.
            if other:
                size_name = other if isinstance(other, str) else 'small'
                if size_name == 'large':
                    line = '_' * max(self.settings["write_in"]["small_underscore_count"], int((widths[1] - .16) * 72 / (size * .56)))
                    label += '\n' + '\n'.join([line] * self.settings["write_in"]["large_lines"])
                elif size_name == 'medium':
                    last = label.split('\n')[-1]
                    remaining = (widths[1] - .16) * 72 - text_width(last + ' ', size)
                    if remaining < size * .56:
                        label += '\n'
                        remaining = (widths[1] - .16) * 72
                    else:
                        label += ' '
                    label += '_' * max(1, int(remaining / (size * .56)))
                else:
                    label += ' ' + '_' * self.settings["write_in"]["small_underscore_count"]
            self.cell_text(label_cell, label, size=size)
            if other == 'large':
                label_cell.paragraphs[0].paragraph_format.line_spacing = self.settings["write_in"]["large_line_spacing"]
            label_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    def text_entry_lines(self, q, selector):
        if selector == "SL":
            return 1
        configured = q.get("Configuration", {})
        if not isinstance(configured, dict):
            configured = {}
        height = configured.get("InputHeight")
        try:
            # Qualtrics records editor box heights in pixels. Approximate their
            # relative size while retaining enough room for handwriting.
            height_lines = math.ceil(float(height) / 28)
        except (ValueError, TypeError):
            height_lines = 0
        base = max(self.text_lines, 5) if selector in {"ESTB", "ESSAY"} else self.text_lines
        return min(12, max(base, height_lines))

    def grid(self, rows, columns, multiple=False, available_height=640):
        if not columns:
            self.warnings.append("Matrix had no scale points; printed a blank response column for review.")
            columns = [""]
        rows = [flat(label) for label in rows]
        labels = [flat(label) for label in columns]
        cfg, response = self.settings["matrix"], self.settings["responses"]
        rotate = cfg["header_rotation"] == "always" or (cfg["header_rotation"] == "auto" and len(columns) >= cfg["rotate_at_columns"])
        base_width = cfg["statement_column_width_inches"] or (3.0 if len(columns) >= 9 else 3.2 if rotate else 2.55) * self.width / 6.9
        preferred_height = (cfg["preferred_row_height_inches"] * 72 if cfg["preferred_row_height_inches"] is not None else
                            32.4 if len(rows) <= 8 and max((len(label) for label in rows), default=0) > 55 else 25.2)
        preferred_height = max(cfg["minimum_row_height_inches"] * 72, preferred_height)
        max_body_height = cfg["maximum_row_height_inches"] * 72
        max_header_height = cfg["header_max_height_inches"] * 72 if cfg["header_max_height_inches"] is not None else (115.2 if len(columns) >= 9 else 126)
        body_size = cfg["statement_font_size"] or response["font_size"]
        min_body_size = cfg["minimum_statement_font_size"] or response["minimum_font_size"]
        spacing = cfg["line_spacing"]
        padding = cfg["cell_padding_inches"] or 0
        pad_pt = padding * 144
        min_header_size = cfg["minimum_header_font_size"]
        header_sizes = sizes_between(cfg["header_font_size"], min_header_size)
        def row_height(label, first, size):
            return estimated_lines(label, first + .16 - padding * 2, size / 1.12) * size * spacing + max(2, pad_pt)
        minimum_scale_width = min(.24, self.width / (2 * len(columns)))
        widths_to_try = [base_width] + [v * self.width / 6.9 for v in (3.1, 3.4, 3.8, 2.25, 2.0)]
        widths_to_try = [min(self.width - minimum_scale_width * len(columns), max(.5, w)) for w in widths_to_try]
        layout = None
        # Find a readable layout within one page, including its question prompt.
        # Header labels shrink to the configured minimum, then wrap.
        for statement_size in sizes_between(body_size, min_body_size):
            for first in widths_to_try:
                scale_width = (self.width - first) / len(columns)
                if scale_width < minimum_scale_width - .0001:
                    continue
                body_height = max(preferred_height, response["symbol_size"] * spacing + max(2, pad_pt),
                                  max((row_height(label, first, statement_size) for label in rows), default=14.4))
                if body_height > max_body_height and statement_size > min_body_size:
                    continue
                header_limit = min(max_header_height, available_height - body_height * len(rows))
                if header_limit < (36 if rotate else 18):
                    continue
                for scale_size in header_sizes:
                    if rotate:
                        required = max((text_width(label, scale_size, bold=cfg["header_bold"]) * 1.12 + max(12, pad_pt)
                                        for label in labels), default=18)
                        if required <= header_limit:
                            header_height = max(36, required)
                        elif scale_size == min_header_size:
                            header_height = header_limit
                            wrapped_lines = max((estimated_lines(label, max(.2, (header_height - 12) / 72) + .16, scale_size)
                                                 for label in labels), default=1)
                            if wrapped_lines * scale_size * max(1.2, spacing) + max(2, pad_pt) > scale_width * 72 - pad_pt:
                                continue
                        else:
                            continue
                    else:
                        header_height = max(18, max((estimated_lines(label, scale_width + .16 - padding * 2, scale_size) * scale_size * max(1.25, spacing) + max(4, pad_pt)
                                                    for label in labels), default=18))
                        if header_height > header_limit:
                            continue
                    layout = (first, scale_width, statement_size, scale_size, body_height, header_height)
                    break
                if layout:
                    break
            if layout:
                break
        if layout is None:
            # Formatting preferences must never stop conversion. Split very
            # wide scales into editable panels before compromising readability.
            if len(columns) > 12:
                self.warnings.append('Wide matrix divided into column panels; every statement and scale point is retained.')
                for offset in range(0, len(columns), 10):
                    self.grid(rows, labels[offset:offset + 10], multiple, available_height)
                return
            # Keep the preferred text direction when relaxing layout limits.
            # Narrow scale columns need a taller rotated header, not a switch
            # to horizontal labels that wrap into fragments of words.
            first = min(self.width - minimum_scale_width * len(columns), max(.5, base_width if cfg["statement_column_width_inches"] is not None else (3.8 if rows else 2.55) * self.width / 6.9))
            scale_width = (self.width - first) / len(columns)
            scale_size = min_header_size
            statement_size = body_size
            def relaxed_body_height(size):
                return max(preferred_height, response["symbol_size"] * spacing + max(2, pad_pt), max((row_height(label, first, size) for label in rows), default=14.4))
            while statement_size > min_body_size and relaxed_body_height(statement_size) > max_body_height:
                statement_size = max(min_body_size, statement_size - 1)
            body_height = relaxed_body_height(statement_size)
            if rotate:
                required = max((text_width(label, scale_size, bold=cfg["header_bold"]) * 1.12 + max(12, pad_pt)
                                for label in labels), default=36)
                header_height = max(36, required)
                # Prefer 1.75 inches if wrapping can fit across the column.
                for candidate in sorted(set([max_header_height, 126, 144, 162, 180, 216, 288, 360])):
                    lines = max((estimated_lines(label, (candidate - 12) / 72 + .16, scale_size)
                                 for label in labels), default=1)
                    if lines * scale_size * max(1.2, spacing) + max(2, pad_pt) <= scale_width * 72 - pad_pt:
                        header_height = min(header_height, candidate)
                        break
            else:
                header_height = max(18, max((estimated_lines(label, scale_width + .16 - padding * 2, scale_size) * scale_size * max(1.25, spacing) + max(4, pad_pt)
                                            for label in labels), default=18))
            layout = (first, scale_width, statement_size, scale_size, body_height, header_height)
            relaxed = True
            self.warnings.append('Matrix sizing preferences were relaxed to retain readable scale labels and all statements; it may continue onto additional pages.')
        else:
            relaxed = False
        first, scale_width, statement_size, scale_size, body_height, header_height = layout
        widths = [first] + [scale_width] * len(columns)
        table = self.doc.add_table(rows=len(rows) + 1, cols=len(widths))
        if relaxed and header_height + body_height * len(rows) > available_height:
            self.flowing_tables.add(table._tbl)
        table.autofit = False
        for column, width in zip(table.columns, widths):
            column.width = Inches(width)
        margins = OxmlElement('w:tblCellMar')
        for edge, amount in [('top', 14), ('bottom', 14), ('left', 14), ('right', 14)]:
            node = OxmlElement('w:' + edge)
            node.set(qn('w:w'), str(amount))
            node.set(qn('w:type'), 'dxa')
            margins.append(node)
        table._tbl.tblPr.append(margins)
        for row in table.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Inches(width)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                paragraph_zero(cell.paragraphs[0])
        for i, label in enumerate([''] + labels):
            cell = table.cell(0, i)
            self.cell_text(cell, label, cfg["header_bold"], scale_size)
            cell.paragraphs[0].paragraph_format.line_spacing = spacing
            shade(cell, cfg["header_color"].lstrip("#"))
            if i:
                if rotate:
                    direction = OxmlElement('w:textDirection')
                    direction.set(qn('w:val'), 'btLr')  # 90 degrees counterclockwise
                    cell._tc.get_or_add_tcPr().append(direction)
                    # Rotation swaps the physical alignment axes: left means
                    # bottom along the text axis; center means horizontal center.
                    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                else:
                    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if not relaxed or header_height <= 300:
            repeat_header(table.rows[0])
        for n, label in enumerate(rows, 1):
            cells = table.rows[n].cells
            self.cell_text(cells[0], label, size=statement_size)
            cells[0].paragraphs[0].paragraph_format.line_spacing = Pt(statement_size * spacing)
            for cell in cells[1:]:
                self.cell_text(cell, response["multiple_mark"] if multiple else response["single_mark"], size=response["symbol_size"])
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                cell.paragraphs[0].paragraph_format.line_spacing = Pt(response["symbol_size"] * spacing)
        self.decorate_table(table, matrix=True)
        for i, row in enumerate(table.rows):
            # Allow an exceptionally tall row to flow rather than clip its text.
            if not relaxed or (header_height if i == 0 else body_height) <= 600:
                no_split(row)
        table.rows[0].height = Pt(header_height)
        table.rows[0].height_rule = WD_ROW_HEIGHT_RULE.EXACTLY if rotate and not relaxed else WD_ROW_HEIGHT_RULE.AT_LEAST
        for row in table.rows[1:]:
            row.height = Pt(body_height)
            row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST if relaxed else WD_ROW_HEIGHT_RULE.EXACTLY
        if body_height > max_body_height:
            self.warnings.append('Matrix response rows exceed the configured maximum to preserve their text.')

    def question(self, qid, q, start_new_page=False, follows_block_heading=False):
        self.settings = merge_config(self.base_config, self.base_config["question_overrides"].get(qid, {}))
        try:
            self._question(qid, q, start_new_page, follows_block_heading)
        finally:
            self.settings = self.base_config

    def _question(self, qid, q, start_new_page=False, follows_block_heading=False):
        if not isinstance(q, dict):
            return
        qtype = str(q.get("QuestionType", ""))
        selector = str(q.get("Selector", ""))
        prompt = clean(q.get("QuestionText"))
        cfg = self.settings["questions"]
        font = self.settings["fonts"]
        response = self.settings["responses"]
        write = self.settings["write_in"]
        pagination = self.settings["pagination"]
        start_new_page = start_new_page or cfg["page_break_before"]
        start = len(list(self.doc.element.body)) - 1
        first_gap = None
        for _ in range(cfg["blank_lines_before"]):
            gap = self.doc.add_paragraph()
            if first_gap is None:
                first_gap = gap
            gap.paragraph_format.space_before = Pt(0)
            gap.paragraph_format.space_after = Pt(0)
            gap.paragraph_format.line_spacing = Pt(cfg["font_size"])
            gap.paragraph_format.keep_with_next = True
        if first_gap is not None:
            first_gap.paragraph_format.page_break_before = start_new_page
        self.number += 1
        tag = clean(q.get("DataExportTag")) if cfg["numbering"] == "survey" else str(self.number)
        tag = tag or str(self.number)
        if qtype == "DB":
            self.number -= 1
            if prompt:
                p = self.doc.add_paragraph()
                rich_runs(p, prompt, cfg["font_size"], font=font["family"], preserve=font["preserve_rich_text"], force_bold=cfg["bold"])
                p.paragraph_format.line_spacing = cfg["line_spacing"]
                p.paragraph_format.space_after = Pt(cfg["space_after_pt"])
                if first_gap is None: p.paragraph_format.page_break_before = start_new_page
                p.paragraph_format.keep_together = True
                p.paragraph_format.keep_with_next = False
            return
        choices = ordered(q.get("Choices"), q.get("ChoiceOrder"))
        answers = ordered(q.get("Answers"), q.get("AnswerOrder"))
        short_write_in = write["inline_short_answers"] and ((qtype == 'TE' and selector == 'SL') or
                          (qtype == 'MC' and selector == 'DL' and len(choices) >= response["dropdown_write_in_at"])) and len(prompt.split('\n')[0]) <= write["inline_prompt_max_chars"]
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.keep_with_next = True
        if first_gap is None:
            first_gap = p
            p.paragraph_format.page_break_before = start_new_page
        p.paragraph_format.line_spacing = cfg["line_spacing"]
        p.paragraph_format.space_after = Pt(cfg["space_after_pt"])
        if cfg["numbering"] != "off": p.add_run(f"{tag}.  ").bold = True
        rich_runs(p, inline_blank(prompt, write["small_underscore_count"]) if short_write_in else prompt,
                  cfg["font_size"], font=font["family"], preserve=font["preserve_rich_text"], force_bold=cfg["bold"])
        for run in p.runs:
            run.font.name = font["family"]
            run.font.size = Pt(cfg["font_size"])
        labels = [(clean(c), other_size(c)) for _, c in choices]
        sec = self.doc.sections[0]
        page_height = (sec.page_height - sec.top_margin - sec.bottom_margin) / 914400 * 72
        prompt_height = estimated_lines((str(tag) + '. ' if cfg["numbering"] != "off" else '') + prompt, self.width, cfg["font_size"]) * cfg["font_size"] * (cfg["line_spacing"] or 1.25) + cfg["space_after_pt"]
        # A conservative full-page budget: Word handles the final placement of
        # the kept-together question within the current or next page.
        response_budget = page_height - cfg["blank_lines_before"] * cfg["font_size"] - prompt_height - 48
        if short_write_in:
            pass
        elif qtype == "MC":
            if selector == "DL" and len(labels) >= response["dropdown_write_in_at"]:
                self.lines(1, indent_inches=response["indent_inches"])
            else:
                self.options(labels, selector.startswith("MA"),
                             rating=not selector.startswith("MA") and is_rating_scale(q, labels))
        elif qtype == "Matrix":
            multiple = selector.startswith("MA") or q.get("SubSelector") in {"MultipleAnswer", "MultiAnswer"}
            self.grid([x for x, _ in labels], [clean(a) for _, a in answers], multiple,
                      available_height=response_budget)
            # Word's keep-with-next behavior varies around large tables.
            # Explicitly start large matrix questions on a fresh page, above
            # their two blank lines and prompt rather than between prompt/table.
            from docx.table import Table
            matrices = [Table(element, self.doc) for element in list(self.doc.element.body)[start:]
                        if element.tag == qn('w:tbl')]
            table_height = sum(sum(row.height.pt if row.height is not None else 18
                                   for row in table.rows) for table in matrices)
            if pagination["new_page_for_large_matrices"] and not follows_block_heading and (len(labels) >= pagination["large_matrix_min_rows"] or len(answers) >= pagination["large_matrix_min_columns"] or table_height >= pagination["large_matrix_min_height_inches"] * 72):
                first_gap.paragraph_format.page_break_before = True
        elif qtype == "TE":
            if selector == "FORM" and choices:
                remaining = self.remaining_page_height(start) if not start_new_page else page_height
                current_budget = remaining - cfg["blank_lines_before"] * cfg["font_size"] - prompt_height - 24
                self.fields([label for label, _ in labels], available_height=response_budget,
                            current_page_height=current_budget)
            else:
                self.lines(self.text_entry_lines(q, selector))
        elif qtype in {"RankOrder", "RO"}:
            if labels:
                self.fields([label for label, _ in labels], answer_first=True)
            else:
                self.lines(1)
        elif qtype == "CS":
            if labels:
                self.fields([label for label, _ in labels])
            else:
                self.lines(1)
        elif qtype == "Slider":
            if labels:
                self.fields([label for label, _ in labels])
            else:
                self.lines(1)
            self.warnings.append(f"{qid}: check slider endpoints and increments against the online survey.")
        elif qtype == "SBS":
            if labels:
                self.fields([label for label, _ in labels])
            else:
                self.lines()
            self.warnings.append(f"{qid}: side by side columns need manual verification.")
        else:
            if labels:
                self.options(labels)
            else:
                self.lines()
            self.warnings.append(f"{qid}: unsupported question type {qtype or 'unknown'} / {selector or 'unknown'}.")
        self.keep_question(start)


def flow_nodes(node):
    if isinstance(node, list):
        return node
    if isinstance(node, dict):
        return node.get("Flow", [])
    return []


def convert(qsf, output, text_lines=None, include_ids=False, config=None):
    with open(qsf, encoding="utf-8-sig") as f:
        data = json.load(f)
    elements = data.get("SurveyElements")
    if not isinstance(elements, list):
        raise ValueError("No SurveyElements array found; is this a QSF survey export?")
    by_type = {}
    for e in elements:
        by_type.setdefault(e.get("Element"), []).append(e.get("Payload", {}))
    questions = {q.get("QuestionID"): q for q in by_type.get("SQ", []) if isinstance(q, dict)}
    blocks = {}
    for payload in by_type.get("BL", []):
        items = payload if isinstance(payload, list) else [payload]
        for block in items:
            if isinstance(block, dict) and block.get("ID"):
                blocks[block["ID"]] = block
    title = clean(data.get("SurveyEntry", {}).get("SurveyName", "")) or Path(qsf).stem
    writer = Writer(title, text_lines, include_ids, config)
    for qid in writer.base_config["question_overrides"]:
        if qid not in questions:
            writer.warnings.append(f"Question override {qid} does not match a question in this QSF.")
    visited = set()
    printed_questions = set()
    printed_blocks = set()
    trash_questions = set()

    def is_trash(block):
        return (str(block.get("Type", "")).lower() == "trash" or
                "trash" in clean(block.get("Description", "")).lower())

    for block in blocks.values():
        if is_trash(block):
            trash_questions.update(item.get("QuestionID") for item in block.get("BlockElements", [])
                                   if isinstance(item, dict) and item.get("Type") == "Question")

    def emit_block(block_id):
        block = blocks.get(block_id)
        if not block:
            writer.warnings.append(f"Flow refers to missing block {block_id}.")
            return
        if block_id in visited:
            return
        visited.add(block_id)
        if is_trash(block):
            return
        # Empty blocks and repeated references do not create blank pages.
        pending_questions = [item.get("QuestionID") for item in block.get("BlockElements", [])
                             if isinstance(item, dict) and item.get("Type") == "Question"
                             and item.get("QuestionID") in questions
                             and item.get("QuestionID") not in printed_questions]
        if not pending_questions:
            return
        new_block_page = bool(printed_blocks) and writer.base_config["pagination"]["new_page_per_block"]
        printed_blocks.add(block_id)
        description = clean(block.get("Description", ""))
        heading = None
        if writer.base_config["sections"]["show_block_names"] and description and description.lower() not in {"default question block", "trash / unused questions"}:
            heading = writer.doc.add_paragraph(description, style="Heading 1")
            heading.paragraph_format.keep_with_next = True
            heading.paragraph_format.page_break_before = new_block_page
        first_question = True
        for item in block.get("BlockElements", []):
            kind = item.get("Type")
            if kind == "Question":
                qid = item.get("QuestionID")
                if qid in questions and qid not in printed_questions:
                    writer.question(qid, questions[qid],
                                    start_new_page=first_question and new_block_page and heading is None,
                                    follows_block_heading=first_question and heading is not None)
                    printed_questions.add(qid)
                    first_question = False
                elif qid in printed_questions:
                    continue
                else:
                    writer.warnings.append(f"Missing question {qid} in block {block_id}.")
            # Online page breaks do not constrain the paper layout. Word's
            # keep-with-next chains paginate complete questions instead.

    def walk(nodes):
        for node in flow_nodes(nodes):
            if not isinstance(node, dict):
                continue
            kind = node.get("Type")
            if kind in {"Block", "Standard"}:
                emit_block(node.get("ID"))
            elif kind in {"Branch", "Randomizer", "Group", "LoopAndMerge"}:
                walk(node.get("Flow", []))
            elif kind in {"EmbeddedData", "EndSurvey", "WebService", "ReferenceSurvey", "Authenticator"}:
                continue
            elif node.get("Flow"):
                walk(node.get("Flow"))

    flows = by_type.get("FL", [])
    for flow in flows:
        walk(flow)
    # Some exports omit a block from flow, or place a question outside blocks.
    # Preserve every non-trash question in the paper backup rather than silently dropping it.
    unvisited = [bid for bid, block in blocks.items() if bid not in visited and not is_trash(block)]
    for block_id in unvisited:
        emit_block(block_id)
    orphaned = [qid for qid in questions if qid not in printed_questions and qid not in trash_questions]
    for qid in orphaned:
        writer.question(qid, questions[qid])
        printed_questions.add(qid)
    if unvisited or orphaned:
        writer.warnings.append(f"Included {len(unvisited)} blocks absent from flow and {len(orphaned)} questions outside blocks.")
    if not questions:
        raise ValueError("No questions found in this QSF.")
    if not writer.number:
        writer.warnings.append("No response questions were printed; inspect the exported QSF.")
    writer.doc.save(output)
    return writer.number, writer.warnings


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("qsf", type=Path, nargs="?", help="Qualtrics .qsf export")
    ap.add_argument("-o", "--output", type=Path, help="Output .docx path (default: <input>-paper.docx)")
    ap.add_argument("--config", type=Path, help="Main JSON defaults (automatically uses config.json beside the script if omitted)")
    ap.add_argument("--survey-config", type=Path, help="Survey-specific JSON overrides layered over the main config")
    ap.add_argument("--set", action="append", default=[], metavar="SECTION.KEY=VALUE", help="Override any setting; repeat as needed. Values may be JSON or plain strings.")
    ap.add_argument("--write-default-config", type=Path, metavar="PATH", help="Write all built-in defaults to JSON and exit")
    ap.add_argument("--list-questions", action="store_true", help="List QSF question IDs, export tags, and prompts, then exit")
    ap.add_argument("--text-lines", type=int, help="Lines for long text answers")
    ap.add_argument("--use-export-tags", action="store_true", help="Number questions with Qualtrics export tags")
    ap.add_argument("--numbering", choices=["off", "sequential", "survey"])
    ap.add_argument("--font", help="Font family")
    ap.add_argument("--question-font-size", type=float)
    ap.add_argument("--response-font-size", type=float)
    ap.add_argument("--matrix-header-font-size", type=float)
    ap.add_argument("--matrix-row-height", type=float, metavar="INCHES")
    ap.add_argument("--header-rotation", choices=["auto", "always", "never"])
    ap.add_argument("--question-line-spacing", type=float)
    ap.add_argument("--table-line-spacing", type=float)
    ap.add_argument("--blank-lines-before", type=int)
    def toggle(flag, help_text):
        group = ap.add_mutually_exclusive_group()
        dest = flag.replace("-", "_")
        group.add_argument("--" + flag, dest=dest, action="store_true", help=help_text)
        group.add_argument("--no-" + flag, dest=dest, action="store_false")
        ap.set_defaults(**{dest: None})
    toggle("bold-questions", "Apply bold to all question text")
    toggle("matrix-banding", "Shade alternating matrix response rows")
    toggle("matrix-row-borders", "Draw bottom borders on matrix response rows")
    ap.add_argument("--matrix-row-border-weight", type=float, metavar="POINTS")
    toggle("new-page-per-block", "Start each later block on a fresh page")
    toggle("show-block-names", "Show survey block descriptions as section headings")
    toggle("keep-questions-together", "Keep question prompts and response tables together where possible")
    args = ap.parse_args()
    try:
        if args.write_default_config:
            args.write_default_config.write_text(json.dumps(DEFAULT_CONFIG, indent=2) + "\n", encoding="utf-8")
            print(f"Wrote {args.write_default_config}")
            return
        if args.qsf is None: ap.error("Provide a QSF path, or use --write-default-config PATH.")
        if args.list_questions:
            data = json.loads(args.qsf.read_text(encoding="utf-8-sig"))
            for element in data.get("SurveyElements", []):
                if element.get("Element") == "SQ" and isinstance(element.get("Payload"), dict):
                    q = element["Payload"]
                    print(f"{q.get('QuestionID', '')}\t{clean(q.get('DataExportTag', ''))}\t{flat(clean(q.get('QuestionText', '')))}")
            return
        config, config_warnings = load_config_layers(args.config, args.survey_config)
        cli = {}
        shortcuts = {"font": "fonts.family", "question_font_size": "questions.font_size",
                     "response_font_size": "responses.font_size", "matrix_header_font_size": "matrix.header_font_size",
                     "matrix_row_height": "matrix.preferred_row_height_inches", "header_rotation": "matrix.header_rotation",
                     "numbering": "questions.numbering", "text_lines": "write_in.text_lines",
                     "bold_questions": "questions.bold", "matrix_banding": "matrix.banded_rows",
                     "matrix_row_borders": "matrix.row_bottom_border.enabled",
                     "matrix_row_border_weight": "matrix.row_bottom_border.weight_pt",
                     "question_line_spacing": "questions.line_spacing", "table_line_spacing": "tables.line_spacing",
                     "blank_lines_before": "questions.blank_lines_before",
                     "new_page_per_block": "pagination.new_page_per_block",
                     "show_block_names": "sections.show_block_names",
                     "keep_questions_together": "pagination.keep_questions_together"}
        for arg, key in shortcuts.items():
            value = getattr(args, arg)
            if value is not None: config_assignment(cli, key + "=" + json.dumps(value))
        if args.use_export_tags:
            if args.numbering is not None and args.numbering != "survey":
                raise ValueError("--use-export-tags conflicts with --numbering; use --numbering survey.")
            config_assignment(cli, "questions.numbering=survey")
        for assignment in args.set: config_assignment(cli, assignment)
        resolved = resolve_config(config, cli)
        output = args.output or args.qsf.with_name(args.qsf.stem + "-paper.docx")
        count, warnings = convert(args.qsf, output, config=resolved)
        warnings = config_warnings + warnings
    except (OSError, ValueError, json.JSONDecodeError) as exc:
        ap.exit(1, f"Error: {exc}\n")
    print(f"Wrote {output} ({count} response questions).")
    if warnings:
        print("Review these items against the original survey before printing:", file=sys.stderr)
        for warning in dict.fromkeys(warnings):
            print("- " + warning, file=sys.stderr)


if __name__ == "__main__":
    main()
