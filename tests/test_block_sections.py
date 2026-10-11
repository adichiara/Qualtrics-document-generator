"""Check independent block headings and page breaks through config and CLI."""
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import qsf_to_paper as formatter


def fixture():
    return {"SurveyEntry": {"SurveyName": "Block options"}, "SurveyElements": [
        *[{"Element": "SQ", "Payload": {
            "QuestionID": "QID{}".format(i), "QuestionType": "MC", "Selector": "SAVR",
            "QuestionText": "Prompt {}".format(i), "Choices": {"1": {"Display": "Yes"}, "2": {"Display": "No"}}
        }} for i in (1, 2)],
        {"Element": "BL", "Payload": [
            {"ID": "BL1", "Description": "First section", "BlockElements": [{"Type": "Question", "QuestionID": "QID1"}]},
            {"ID": "EMPTY", "Description": "Empty section", "BlockElements": []},
            {"ID": "BL2", "Description": "Second section", "BlockElements": [{"Type": "Question", "QuestionID": "QID2"}]}
        ]},
        {"Element": "FL", "Payload": {"Flow": [{"Type": "Block", "ID": bid}
                                                   for bid in ("BL1", "EMPTY", "BL2", "BL2")]}}
    ]}


class BlockSectionTests(unittest.TestCase):
    def convert(self, data, config=None):
        with tempfile.TemporaryDirectory() as folder:
            qsf, output = Path(folder) / "survey.qsf", Path(folder) / "survey.docx"
            qsf.write_text(json.dumps(data))
            count, warnings = formatter.convert(qsf, output, config=config)
            self.assertEqual(count, 2)
            self.assertFalse(warnings)
            return Document(output)

    def test_defaults_and_independent_combinations(self):
        self.assertFalse(formatter.DEFAULT_CONFIG["sections"]["show_block_names"])
        self.assertFalse(formatter.DEFAULT_CONFIG["pagination"]["new_page_per_block"])
        for headings in (False, True):
            for page_breaks in (False, True):
                with self.subTest(headings=headings, page_breaks=page_breaks):
                    doc = self.convert(fixture(), {"sections": {"show_block_names": headings},
                                                 "pagination": {"new_page_per_block": page_breaks}})
                    names = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
                    self.assertEqual(names, ["First section", "Second section"] if headings else [])
                    breaks = [p for p in doc.paragraphs if p.paragraph_format.page_break_before]
                    self.assertEqual(len(breaks), int(page_breaks))
                    if breaks:
                        self.assertEqual(breaks[0].text, "Second section" if headings else "")
                    self.assertEqual(sum("Prompt 2" in p.text for p in doc.paragraphs), 1)

    def test_generic_heading_and_explicit_question_break(self):
        data = fixture()
        data["SurveyElements"][2]["Payload"][0]["Description"] = "Default Question Block"
        doc = self.convert(data, {"sections": {"show_block_names": True},
                                  "pagination": {"new_page_per_block": True}})
        self.assertEqual([p.text for p in doc.paragraphs if p.style.name == "Heading 1"], ["Second section"])
        doc = self.convert(fixture(), {"question_overrides": {"QID2": {"questions": {"page_break_before": True}}}})
        self.assertEqual(sum(bool(p.paragraph_format.page_break_before) for p in doc.paragraphs), 1)

    def test_cli_flags_and_config_priority(self):
        with tempfile.TemporaryDirectory() as folder:
            folder = Path(folder)
            qsf, output, main, survey = [folder / name for name in ("survey.qsf", "survey.docx", "main.json", "survey.json")]
            qsf.write_text(json.dumps(fixture()))
            main.write_text(json.dumps({"sections": {"show_block_names": True}, "pagination": {"new_page_per_block": True}}))
            survey.write_text(json.dumps({"sections": {"show_block_names": False}, "pagination": {"new_page_per_block": False}}))
            for flags, expected in (([], False), (["--show-block-names", "--new-page-per-block"], True),
                                    (["--show-block-names", "--new-page-per-block", "--set", "sections.show_block_names=false", "--set", "pagination.new_page_per_block=false"], False)):
                result = subprocess.run([sys.executable, str(ROOT / "qsf_to_paper.py"), str(qsf), "-o", str(output),
                                         "--config", str(main), "--survey-config", str(survey), *flags], capture_output=True, text=True)
                self.assertEqual(result.returncode, 0, result.stderr)
                doc = Document(output)
                self.assertEqual(any(p.style.name == "Heading 1" for p in doc.paragraphs), expected)
                self.assertEqual(any(p.paragraph_format.page_break_before for p in doc.paragraphs), expected)


if __name__ == "__main__":
    unittest.main()
