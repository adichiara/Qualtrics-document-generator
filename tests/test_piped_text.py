"""Verify readable/hidden piped text, rich-text spans, and logic integration."""
import copy
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import qsf_to_paper as formatter
import qsf_logic_to_paper as adapter
from test_logic import fixture, question


PIPE = "${e://Field/Issued_Trouser_Size}"


def document_text(doc):
    return "\n".join([p.text for p in doc.paragraphs] +
                     [cell.text for table in doc.tables for row in table.rows for cell in row.cells])


class PipedTextTests(unittest.TestCase):
    def test_readable_tokens_and_removal(self):
        cases = {
            PIPE: "{Issued_Trouser_Size}",
            "${q://QID12/ChoiceGroup/SelectedChoices}": "{QID12/ChoiceGroup/SelectedChoices}",
            "${q://QID12/ChoiceTextEntryValue/3}": "{QID12/ChoiceTextEntryValue/3}",
            "${lm://Field/1}": "{lm: Field/1}",
            "${date://CurrentDate/DM}": "{date: CurrentDate/DM}",
            "${unrecognized://a/b/c}": "{unrecognized: a/b/c}",
            "${CustomToken}": "{CustomToken}"
        }
        for raw, expected in cases.items():
            with self.subTest(raw=raw):
                self.assertEqual(formatter.format_piped_text(raw), expected)
                self.assertEqual(formatter.format_piped_text(raw, False), "")
        self.assertEqual(formatter.format_piped_text("Size " + PIPE + " supplied", False), "Size supplied")
        self.assertEqual(formatter.format_piped_text(PIPE + " / " + PIPE), "{Issued_Trouser_Size} / {Issued_Trouser_Size}")
        self.assertEqual(formatter.format_piped_text("Ordinary {braces} and $5"), "Ordinary {braces} and $5")
        self.assertEqual(formatter.format_piped_text("Incomplete ${e://Field/Name"), "Incomplete ${e://Field/Name")

    def test_rich_text_and_codes_split_across_runs_or_entities(self):
        value = formatter.clean('<i>Before</i> <b>${e://</b><u>Field/Issued_Trouser_Size}</u> <span style="color: red">after</span>')
        result = formatter.format_piped_text(value)
        self.assertEqual(result, "Before {Issued_Trouser_Size} after")
        by_text = {text: style for text, style in result.spans}
        self.assertTrue(by_text["Before"]["italic"])
        self.assertTrue(by_text["{Issued_Trouser_Size}"]["bold"])
        self.assertEqual(by_text["after"]["color"], "FF0000")
        result = formatter.format_piped_text(value, False)
        self.assertEqual(result, "Before after")
        self.assertTrue(any(text == "Before" and style.get("italic") for text, style in result.spans))
        self.assertEqual(formatter.paper_text('&#36;{e://Field/Issued_Trouser_Size}', formatter.DEFAULT_CONFIG), "{Issued_Trouser_Size}")

    def test_questions_responses_matrix_and_form_labels(self):
        for show in (True, False):
            writer = formatter.Writer("Piped-text checks", config={"piped_text": {"show": show}})
            writer.question("QID1", {"QuestionType": "MC", "Selector": "SAVR", "QuestionText": "Rate " + PIPE,
                                     "Choices": {"1": {"Display": "Issued " + PIPE}, "2": {"Display": "Other"}}})
            writer.question("QID2", {"QuestionType": "Matrix", "Selector": "Likert", "QuestionText": "Matrix " + PIPE,
                                     "Choices": {"1": {"Display": "Statement " + PIPE}},
                                     "Answers": {"1": {"Display": "Scale " + PIPE}, "2": {"Display": "Other"}}})
            writer.question("QID3", {"QuestionType": "TE", "Selector": "FORM", "QuestionText": "Details",
                                     "Choices": {"1": {"Display": "Field " + PIPE}}})
            writer.question("QID4", {"QuestionType": "DB", "QuestionText": "Information " + PIPE})
            text = document_text(writer.doc)
            self.assertNotIn("${", text)
            self.assertEqual("{Issued_Trouser_Size}" in text, show)
            self.assertEqual(writer.number, 3)
            self.assertEqual(len(writer.doc.tables), 3)
            self.assertEqual(writer.doc.tables[0].cell(0, 1).text, "Issued {Issued_Trouser_Size}" if show else "Issued")
            self.assertEqual(writer.doc.tables[1].cell(0, 1).text, "Scale {Issued_Trouser_Size}" if show else "Scale")
            self.assertEqual(writer.doc.tables[2].cell(0, 0).text, "Field {Issued_Trouser_Size}" if show else "Field")

    def test_defaults_validation_and_question_overrides(self):
        self.assertTrue(formatter.DEFAULT_CONFIG["piped_text"]["show"])
        self.assertEqual(json.loads((ROOT / "config.json").read_text()), formatter.DEFAULT_CONFIG)
        with self.assertRaises(ValueError): formatter.resolve_config({"piped_text": {"show": "false"}})
        writer = formatter.Writer("Overrides", config={"piped_text": {"show": False}, "question_overrides": {
            "QID1": {"piped_text": {"show": True}}}})
        q = {"QuestionType": "TE", "Selector": "SL", "QuestionText": "Confirm " + PIPE}
        writer.question("QID1", q); writer.question("QID2", q)
        prompts = [p.text for p in writer.doc.paragraphs if "Confirm" in p.text]
        self.assertIn("{Issued_Trouser_Size}", prompts[0])
        self.assertNotIn("{Issued_Trouser_Size}", prompts[1])

    def test_logic_anchors_and_choice_instructions_follow_formatted_text(self):
        for show in (True, False):
            data = fixture()
            question(data, 1)["QuestionText"] = "Confirm " + PIPE
            question(data, 1)["Choices"]["2"]["Display"] = "No " + PIPE
            original = copy.deepcopy(data)
            with tempfile.TemporaryDirectory() as folder:
                qsf, out = Path(folder) / "survey.qsf", Path(folder) / "out.docx"
                qsf.write_text(json.dumps(data))
                count, warnings, report = adapter.convert_with_logic(qsf, out, config={"piped_text": {"show": show}})
                text = document_text(Document(out))
                self.assertEqual(count, 4)
                self.assertNotIn("${", text)
                self.assertIn("Skip to Q3.", text)
                self.assertEqual("{Issued_Trouser_Size}" in text, show)
                self.assertTrue(any("Skip to Q3." in cell.text for t in Document(out).tables for row in t.rows for cell in row.cells))
            self.assertEqual(data, original)
        # Eligibility directions may quote a source response containing a pipe.
        data = fixture(); question(data, 1)["Selector"] = "MAVR"
        question(data, 1)["Choices"]["1"]["Display"] = "Yes " + PIPE
        with tempfile.TemporaryDirectory() as folder:
            qsf, out = Path(folder) / "survey.qsf", Path(folder) / "out.docx"
            qsf.write_text(json.dumps(data))
            adapter.convert_with_logic(qsf, out, config={"piped_text": {"show": False}})
            text = document_text(Document(out))
            self.assertIn("Answer only if", text)
            self.assertNotIn("${", text)

    def test_cli_toggle_both_scripts_and_set_precedence(self):
        with tempfile.TemporaryDirectory() as folder:
            folder = Path(folder)
            data = fixture(); question(data, 1)["QuestionText"] = "Confirm " + PIPE
            qsf, out, config = folder / "survey.qsf", folder / "out.docx", folder / "config.json"
            qsf.write_text(json.dumps(data)); config.write_text(json.dumps({"piped_text": {"show": False}}))
            for script in ("qsf_to_paper.py", "qsf_logic_to_paper.py"):
                for flags, expected in (([], False), (["--show-piped-text"], True),
                                        (["--no-show-piped-text"], False),
                                        (["--show-piped-text", "--set", "piped_text.show=false"], False)):
                    result = subprocess.run([sys.executable, str(ROOT / script), str(qsf), "-o", str(out),
                                             "--config", str(config), *flags], capture_output=True, text=True)
                    self.assertEqual(result.returncode, 0, result.stderr)
                    text = document_text(Document(out))
                    self.assertNotIn("${", text)
                    self.assertEqual("{Issued_Trouser_Size}" in text, expected)


if __name__ == "__main__":
    unittest.main()
