"""Verify instruction insertion and numbering without modifying the base formatter."""
import hashlib
import json
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import qsf_logic_to_paper as adapter
from test_logic import fixture, question, choice, condition


class LogicDocumentTests(unittest.TestCase):
    def convert(self, data, config=None):
        with tempfile.TemporaryDirectory() as folder:
            qsf = Path(folder) / 'survey.qsf'
            output = Path(folder) / 'survey.docx'
            qsf.write_text(json.dumps(data), encoding='utf-8')
            count, warnings, report = adapter.convert_with_logic(qsf, output, config=config)
            return Document(output), report

    def test_choice_cell_instruction_and_symbol_separation(self):
        doc, report = self.convert(fixture())
        cells = [cell for table in doc.tables for row in table.rows for cell in row.cells]
        matches = [cell for cell in cells if 'Skip to Q3.' in cell.text]
        self.assertEqual(len(matches), 1)
        self.assertIn('No', matches[0].text)
        self.assertNotIn('◯', matches[0].text)
        self.assertEqual(matches[0].paragraphs[0].paragraph_format.space_after.pt, 0)

    def test_survey_and_per_question_numbering(self):
        config = {'questions': {'numbering': 'survey'},
                  'question_overrides': {'QID3': {'questions': {'numbering': 'sequential'}}}}
        doc, report = self.convert(fixture(), config)
        self.assertEqual(report['question_references']['QID1'], 'Q10')
        self.assertEqual(report['question_references']['QID3'], 'Q3')
        self.assertTrue(any(p.text.startswith('3.  ') for p in doc.paragraphs))
        self.assertTrue(any('Skip to Q3.' in cell.text for t in doc.tables for row in t.rows for cell in row.cells))

    def test_eligibility_precedes_target_question(self):
        data = fixture()
        question(data, 1)['Selector'] = 'MAVR'
        doc, report = self.convert(data)
        paragraphs = [p for p in doc.paragraphs if p.text]
        target = next(i for i, p in enumerate(paragraphs) if p.text.startswith('2.  '))
        self.assertIn('Answer only if', paragraphs[target - 1].text)
        self.assertTrue(paragraphs[target - 1].paragraph_format.keep_with_next)

    def test_real_survey_and_base_script_unchanged(self):
        source = ROOT / 'qsf_to_paper.py'
        before = hashlib.sha256(source.read_bytes()).hexdigest()
        data = json.loads((ROOT / 'examples/Extreme_Cold_Weather_Fuel_Handler_Handwear_Survey.qsf').read_text(encoding='utf-8-sig'))
        doc, report = self.convert(data)
        self.assertEqual(len(report['question_references']), 79)
        combined = '\n'.join(p.text for p in doc.paragraphs) + '\n' + '\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
        self.assertIn('Skip to the HANDWEAR section.', combined)
        self.assertIn('selected at least one response in Q65', combined)
        self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(), before)


if __name__ == '__main__':
    unittest.main()
