"""Behavior checks for formatting configuration (internal QA)."""
import copy
import json
import subprocess
import sys
import tempfile
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import qsf_to_paper as m


def prompt(doc, text):
    return next(p for p in doc.paragraphs if text in p.text)


def matrix_question(label="Matrix prompt"):
    return {"QuestionType": "Matrix", "Selector": "Likert", "QuestionText": label,
            "DataExportTag": "Q42", "Choices": {"1": {"Display": "First statement"}, "2": {"Display": "Second statement"}},
            "Answers": {str(i): {"Display": "Point " + str(i)} for i in range(1, 6)}}


def test_configuration():
    c=m.resolve_config({"questions":{"font_size":14}, "question_overrides":{"QID1":{"questions":{"font_size":16}}}}, {"questions":{"font_size":13}})
    assert c["questions"]["font_size"]==13
    w=m.Writer('Test', config=c)
    q={"QuestionType":"TE","Selector":"SL","QuestionText":"First prompt"}
    w.question('QID1', q); w.question('QID2',dict(q,QuestionText='Second prompt'))
    assert prompt(w.doc,'First prompt').runs[1].font.size.pt==16
    assert prompt(w.doc,'Second prompt').runs[1].font.size.pt==13
    assert m.DEFAULT_CONFIG["questions"]["font_size"]==12
    for numbering in ('off','sequential','survey'):
        w=m.Writer('Test',config={"questions":{"numbering":numbering,"blank_lines_before":0,"bold":True,"font_size":14},"fonts":{"family":"Calibri"}})
        w.question('QID1',dict(q,QuestionText='<b>Bold</b> ordinary prompt',DataExportTag='Q42'),start_new_page=True)
        p=prompt(w.doc,'ordinary prompt')
        assert p.text.startswith({'off':'Bold','sequential':'1.','survey':'Q42.'}[numbering])
        assert all(r.bold and r.font.size.pt==14 and r.font.name=='Calibri' for r in p.runs)
        assert p.paragraph_format.page_break_before
    for band in (False,True):
        for border in (False,True):
            w=m.Writer('Test',config={"matrix":{"banded_rows":band,"row_bottom_border":{"enabled":border,"weight_pt":1.25,"color":"#123456"}}})
            w.question('QID1',matrix_question())
            t=w.doc.tables[0]
            assert (t.cell(2,0)._tc.find('.//'+qn('w:shd')) is not None)==band
            b=t.cell(1,0)._tc.find('.//'+qn('w:bottom'))
            assert (b is not None)==border
            if b is not None: assert b.get(qn('w:sz'))=='10' and b.get(qn('w:color'))=='123456'
    for rotation in ('always','never'):
        w=m.Writer('Test',config={"matrix":{"header_rotation":rotation,"preferred_row_height_inches":.4,"line_spacing":1.2,"statement_column_width_inches":2.5}})
        w.question('QID1',matrix_question())
        t=w.doc.tables[0]
        assert (t.cell(0,1)._tc.find('.//'+qn('w:textDirection')) is not None)==(rotation=='always')
        assert all(abs(r.height.inches-.4)<.002 for r in t.rows[1:])
        assert abs(t.columns[0].width.inches-2.5)<.002
        assert len(set(c.width for c in t.rows[0].cells[1:]))==1
    w=m.Writer('Test',config={"fonts":{"preserve_rich_text":False},"questions":{"line_spacing":1.5},"tables":{"line_spacing":1.25,"banded_rows":True,"row_bottom_border":{"enabled":True},"borders":{"enabled":True}},"responses":{"row_height_inches":.4}})
    w.question('QID1',{'QuestionType':'MC','Selector':'SAVR','QuestionText':'<b>Plain prompt</b>','Choices':{'1':{'Display':'<b>Choice one</b>'},'2':{'Display':'Choice two'}}})
    assert prompt(w.doc,'Plain prompt').runs[1].bold is False
    t=w.doc.tables[0]
    assert all(r.bold is False for r in t.cell(0,1).paragraphs[0].runs if r.text)
    assert t.cell(0,1).paragraphs[0].paragraph_format.line_spacing==1.25
    assert t.cell(1,1)._tc.find('.//'+qn('w:shd')) is not None
    assert t.cell(0,1)._tc.find('.//'+qn('w:bottom')) is not None
    assert abs(t.rows[0].height.inches-.4)<.002
    bad=[{'matrix':{'typo':True}}, {'questions':{'bold':'yes'}}, {'questions':{'font_size':0}}, {'responses':{'font_size':7}}, {'questions':{'numbering':'invalid'}}, {'tables':{'band_color':'orange'}}, {'page':{'margin_left_inches':8}}, {'matrix':{'row_bottom_border':{'weight_pt':0}}}, {'question_overrides':{'QID1':{'page':{'paper_size':'a4'}}}}]
    for patch in bad:
        try: m.resolve_config(patch)
        except ValueError: pass
        else: raise AssertionError(patch)
    with tempfile.TemporaryDirectory() as tmp:
        tmp=Path(tmp)
        cfg=tmp/'config.json';cfg.write_text(json.dumps({'questions':{'font_size':14},'question_overrides':{'QID2':{'questions':{'font_size':16}}}}))
        out=tmp/'out.docx'
        r=subprocess.run([sys.executable,str(ROOT / 'qsf_to_paper.py'),str(ROOT / 'examples/Example_survey.qsf'),'--config',str(cfg),'--question-font-size','13','--set','questions.numbering=off','-o',str(out)],capture_output=True,text=True)
        assert r.returncode==0,r.stderr
        d=Document(out)
        assert prompt(d,'Please rate your satisfaction').runs[0].font.size.pt==16
        assert prompt(d,'Please indicate your sex').runs[0].font.size.pt==13
        r=subprocess.run([sys.executable,str(ROOT / 'qsf_to_paper.py'),str(ROOT / 'examples/Example_survey.qsf'),'--set','questions.font_size=-1','-o',str(out)],capture_output=True,text=True)
        assert r.returncode!=0 and 'questions.font_size' in r.stderr and 'Traceback' not in r.stderr
    w=m.Writer('Test',config={'page':{'paper_size':'a4','orientation':'landscape'}})
    w.question('QID1',matrix_question())
    assert w.width>10
    assert sum(c.width.inches for c in w.doc.tables[0].columns)<=w.width+.01
    w=m.Writer('Test');w.grid(['A longer statement '+str(i) for i in range(65)],['Scale '+str(i) for i in range(16)],available_height=500)
    assert len(w.doc.tables)==2 and w.flowing_tables
    print('Configuration checks passed: precedence/isolation, numbering, styles, independent banding/borders, rotation, spacing, validation, CLI, page width, and oversized matrix fallback.')


if __name__ == '__main__': test_configuration()
