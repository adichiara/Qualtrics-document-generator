"""Verify dropdown decisions, preserved choices/order, geometry, and overrides."""
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import qsf_to_paper as formatter
import qsf_logic_to_paper as adapter


def question(labels, prompt="Choose a response"):
    return {"QuestionType": "MC", "Selector": "DL", "QuestionText": prompt,
            "Choices": {str(i): {"Display": label} for i, label in enumerate(labels, 1)}}


def table_labels(table):
    # The physical row order interleaves columns; respondents read down each.
    return [table.cell(row, column).text for column in range(1, len(table.columns), 3)
            for row in range(len(table.rows)) if table.cell(row, column).text]


class DropdownLayoutTests(unittest.TestCase):
    def writer(self, labels, config=None, prompt="Choose a response"):
        writer = formatter.Writer("Dropdown checks", config=config)
        writer.question("QID1", question(labels, prompt))
        return writer

    def assert_choices(self, labels, config=None, groups=None, prompt="Choose a response"):
        writer = self.writer(labels, config, prompt)
        self.assertEqual(len(writer.doc.tables), 1)
        table = writer.doc.tables[0]
        self.assertEqual(table_labels(table), labels)
        if groups is not None:
            self.assertEqual(len(table.columns), groups * 3 - 1)
        self.assertAlmostEqual(sum(c.width.inches for c in table.columns) + .5, writer.width, places=2)
        self.assertEqual(table._tbl.tblPr.find(qn("w:tblInd")).get(qn("w:w")), "720")
        for column in range(0, len(table.columns), 3):
            for row in range(len(table.rows)):
                if table.cell(row, column + 1).text:
                    self.assertEqual(table.cell(row, column).text, "◯")
        return writer

    def test_short_lists_and_threshold_boundary(self):
        self.assert_choices(["Alpha", "Beta", "Gamma"], groups=1)
        self.assert_choices(["Name {:02d}".format(i) for i in range(49)], groups=3)
        writer = self.writer(["Name {:02d}".format(i) for i in range(50)])
        self.assertFalse(writer.doc.tables)
        self.assertIn("_" * 20, writer.doc.paragraphs[-1].text)

    def test_numeric_units_rank_and_special_endpoints(self):
        sequences = [
            [str(i) for i in range(18, 101)],
            ["{}-{} years".format(i, i + 4) for i in range(0, 300, 5)],
            ["Less than one year"] + ["{} years".format(i) for i in range(1, 51)],
            ["{} cm".format(i) for i in range(130, 211)],
            ["{} ft {} in".format(i // 12, i % 12) for i in range(48, 121)],
            ["{}' {}\"".format(i // 12, i % 12) for i in range(48, 121)],
            ["{} lb".format(i) for i in range(100, 171)],
            [str(i) for i in range(99, 17, -1)] + ["Prefer not to answer"]
        ]
        for labels in sequences:
            with self.subTest(first=labels[0]):
                self.assertEqual(formatter.ordered_choice_kind([(label, False) for label in labels]), "numeric")
                self.assertFalse(self.writer(labels).doc.tables)
        ranks = ["E-{}".format(i) for i in range(1, 10)] + ["W-{}".format(i) for i in range(1, 6)] + ["O-{}".format(i) for i in range(1, 11)]
        self.assert_choices(ranks, groups=3)
        self.assert_choices(["OR-{}".format(i) for i in range(1, 61)], groups=3)
        self.assert_choices([str(i) for i in range(11)], groups=2)
        self.assertFalse(self.writer([str(i) for i in range(12)]).doc.tables)
        self.assert_choices([str(i) for i in range(19)], {"responses": {"dropdown_numeric_write_in_at": 20}}, groups=2)

    def test_names_and_unsorted_identifiers_are_not_ordered(self):
        for labels in (["Person {:02d}".format(i) for i in range(60)],
                       ["{:02d} John Smith".format(i) for i in range(60)],
                       [str(i) for i in range(0, 60, 2)] + [str(i) for i in range(1, 60, 2)]):
            self.assertIsNone(formatter.ordered_choice_kind([(label, False) for label in labels]))
            writer = self.writer(labels)
            self.assertFalse(writer.doc.tables)
        writer = self.writer(["Person {:02d}".format(i) for i in range(60)],
                             {"write_in": {"inline_short_answers": False}})
        self.assertEqual(len(writer.doc.tables[0].columns), 1)
        self.assertEqual(len(writer.doc.tables[0].rows), 1)
        writer = self.writer(["Person {:02d}".format(i) for i in range(60)], prompt="A long prompt " * 15)
        self.assertEqual(len(writer.doc.tables[0].columns), 1)

    def test_mode_column_threshold_and_question_overrides(self):
        names = ["Person {:02d}".format(i) for i in range(60)]
        self.assert_choices(names, {"responses": {"dropdown_mode": "choices", "dropdown_max_columns": 2}}, groups=2)
        self.assert_choices(names, {"responses": {"dropdown_write_in_at": 75}}, groups=3)
        self.assert_choices([str(i) for i in range(60)], {"responses": {"dropdown_mode": "choices", "dropdown_max_columns": 1}}, groups=1)
        writer = self.writer(["Yes", "No"], {"responses": {"dropdown_mode": "write_in"}})
        self.assertFalse(writer.doc.tables)
        writer = self.writer(names, {"question_overrides": {"QID1": {"responses": {"dropdown_mode": "choices"}}}})
        self.assertEqual(table_labels(writer.doc.tables[0]), names)
        writer.question("QID2", question(names))
        self.assertEqual(len(writer.doc.tables), 1)
        # Non-dropdown questions keep their existing two-column limit.
        q = question(names); q["Selector"] = "SAVR"
        writer.question("QID3", q)
        self.assertEqual(len(writer.doc.tables[-1].columns), 5)

    def test_choice_order_labels_and_typography_are_preserved(self):
        labels = [str(i) for i in range(18, 78)]
        q = question(labels)
        q["ChoiceOrder"] = list(reversed(list(q["Choices"])))
        writer = formatter.Writer("Order check", config={"responses": {"dropdown_mode": "choices"}})
        writer.question("QID1", q)
        self.assertEqual(table_labels(writer.doc.tables[0]), list(reversed(labels)))
        table = writer.doc.tables[0]
        self.assertIsNotNone(table.cell(1, 1)._tc.find(".//" + qn("w:shd")))
        sizes = {run.font.size.pt for row in table.rows for cell in row.cells for p in cell.paragraphs
                 for run in p.runs if run.text.isdigit()}
        self.assertEqual(len(sizes), 1)
        self.assertTrue(all(p.paragraph_format.space_before.pt == p.paragraph_format.space_after.pt == 0
                            for row in table.rows for cell in row.cells for p in cell.paragraphs))
        q = question(["A label that is much longer than the allowed column splitting threshold {}".format(i) for i in range(30)])
        writer.question("QID2", q)
        self.assertEqual(len(writer.doc.tables[-1].columns), 2)

    def test_real_dropdowns_keep_grades_and_use_numeric_write_ins(self):
        data = json.loads((ROOT / "examples/FY27_MRE_IOTE.qsf").read_text(encoding="utf-8-sig"))
        writer = formatter.Writer("MRE dropdown checks")
        count = 0
        for element in data["SurveyElements"]:
            q = element.get("Payload", {})
            if element.get("Element") != "SQ" or not isinstance(q, dict) or q.get("Selector") != "DL":
                continue
            before = len(writer.doc.tables)
            writer.question(q["QuestionID"], q)
            expected = [str(formatter.flat(formatter.clean(choice))) for _, choice in formatter.ordered(q["Choices"], q.get("ChoiceOrder"))]
            if q["QuestionID"] == "QID19789337":
                self.assertEqual(len(writer.doc.tables), before + 1)
                self.assertEqual(table_labels(writer.doc.tables[-1]), expected)
            else:
                self.assertEqual(len(writer.doc.tables), before)
                self.assertIn("_" * 20, writer.doc.paragraphs[-1].text)
            count += 1
        self.assertEqual(count, 4)

    def test_config_validation_and_cli(self):
        for patch in ({"dropdown_mode": "unknown"}, {"dropdown_max_columns": 0}, {"dropdown_max_columns": 2.5},
                      {"dropdown_numeric_write_in_at": 0}, {"dropdown_numeric_write_in_at": 12.5}):
            with self.assertRaises(ValueError):
                formatter.resolve_config({"responses": patch})
        self.assertEqual(json.loads((ROOT / "config.json").read_text()), formatter.DEFAULT_CONFIG)
        with tempfile.TemporaryDirectory() as folder:
            out = Path(folder) / "survey.docx"
            result = subprocess.run([sys.executable, str(ROOT / "qsf_to_paper.py"), str(ROOT / "examples/FY27_MRE_IOTE.qsf"),
                                     "-o", str(out), "--set", "question_overrides.QID19789334.responses.dropdown_mode=write_in"], capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stderr)
            doc = Document(out)
            self.assertTrue(any("Please select your age" in p.text and "_" * 20 in p.text for p in doc.paragraphs))

    def test_logic_adapter_uses_tables_and_keeps_skip_instructions(self):
        with tempfile.TemporaryDirectory() as folder:
            out = Path(folder) / "survey.docx"
            count, warnings, report = adapter.convert_with_logic(ROOT / "examples/FY27_MRE_IOTE.qsf", out)
            self.assertEqual(count, 44)
            doc = Document(out)
            ranks = ["E-{}".format(i) for i in range(1, 10)] + ["W-{}".format(i) for i in range(1, 6)] + ["O-{}".format(i) for i in range(1, 11)]
            self.assertTrue(any(table_labels(t) == ranks
                                for t in doc.tables if len(t.columns) == 8))
            self.assertTrue(any("Less than one year" in p.text and "skip to Q9" in p.text for p in doc.paragraphs))


if __name__ == "__main__":
    unittest.main()
