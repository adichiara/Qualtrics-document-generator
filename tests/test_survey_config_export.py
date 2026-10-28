"""Check editable survey config snapshots and conversion round trips."""
import copy
import json
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import qsf_to_paper as formatter
import qsf_logic_to_paper as adapter
from test_logic import fixture, question


class SurveyConfigExportTests(unittest.TestCase):
    def test_full_effective_values_metadata_and_unused_records(self):
        data = fixture(5)
        question(data, 1)["QuestionText"] = '<b>Issued item</b><br>Confirm the size'
        data["SurveyElements"][-2]["Payload"][0]["BlockElements"].pop()
        data["SurveyElements"][-2]["Payload"].append({"ID": "TRASH", "Type": "Trash",
            "Description": "Trash / Unused Questions", "BlockElements": [{"Type": "Question", "QuestionID": "QID5"}]})
        config = formatter.resolve_config({"questions": {"font_size": 13}, "page": {"margin_left_inches": 1},
            "question_overrides": {"QID1": {"responses": {"font_size": 10}},
                                   "QID2": {"_comment": "My handwritten note", "questions": {"font_size": 16}},
                                   "QID404": {"questions": {"bold": True}}}}, {"questions": {"font_size": 14}})
        before = copy.deepcopy(config)
        with tempfile.TemporaryDirectory() as folder:
            source, output = Path(folder) / "survey.qsf", Path(folder) / "settings.json"
            source.write_text(json.dumps(data), encoding="utf-8-sig")
            original = source.read_bytes()
            path, count, warnings = formatter.export_survey_config(source, output, config)
            self.assertEqual(path, output)
            self.assertEqual(count, 5)
            self.assertTrue(any("QID404" in warning for warning in warnings))
            exported = json.loads(output.read_text())
            self.assertEqual(list(exported["question_overrides"]), ["QID1", "QID2", "QID3", "QID4", "QID5"])
            one, two, three = [exported["question_overrides"][f"QID{i}"] for i in (1, 2, 3)]
            self.assertEqual(one["questions"]["font_size"], 14)
            self.assertEqual(one["responses"]["font_size"], 10)
            self.assertEqual(two["questions"]["font_size"], 16)
            self.assertEqual(three["questions"]["font_size"], 14)
            self.assertEqual(one["_comment"], "Q10 | MC/SAVR | Issued item Confirm the size")
            self.assertEqual(two["_comment"], "My handwritten note")
            self.assertEqual(exported["page"]["margin_left_inches"], 1)
            self.assertEqual(set(one), {"_comment", "fonts", "piped_text", "questions", "responses", "tables", "matrix", "write_in", "pagination"})
            self.assertEqual(set(one["fonts"]), {"family", "preserve_rich_text"})
            self.assertNotIn("new_page_per_block", one["pagination"])
            self.assertIsNone(one["matrix"]["preferred_row_height_inches"])
            self.assertEqual(source.read_bytes(), original)
            formatter.resolve_config(exported)
        self.assertEqual(config, before)

    def test_conversion_round_trip_and_editing_one_question(self):
        data = fixture()
        question(data, 1)["QuestionText"] += " ${e://Field/Item}"
        config = {"questions": {"font_size": 13}, "piped_text": {"show": False},
                  "question_overrides": {"QID2": {"questions": {"bold": True}}}}
        with tempfile.TemporaryDirectory() as folder:
            folder = Path(folder)
            source, settings, original, reloaded, edited = [folder / name for name in
                ("survey.qsf", "settings.json", "original.docx", "reloaded.docx", "edited.docx")]
            source.write_text(json.dumps(data))
            formatter.export_survey_config(source, settings, config)
            loaded, warnings = formatter.load_config_layers(survey_path=settings)
            self.assertFalse(warnings)
            for converter in (formatter.convert, adapter.convert_with_logic):
                converter(source, original, config=config)
                converter(source, reloaded, config=loaded)
                self.assertEqual(Document(original).element.xml, Document(reloaded).element.xml)
            exported = json.loads(settings.read_text())
            exported["question_overrides"]["QID3"]["questions"]["font_size"] = 18
            exported["question_overrides"]["QID3"]["_comment"] = "Only this question changes"
            settings.write_text(json.dumps(exported))
            loaded, warnings = formatter.load_config_layers(survey_path=settings)
            formatter.convert(source, edited, config=loaded)
            prompts = [p for p in Document(edited).paragraphs if ".  Question " in p.text]
            sizes = {p.text.split(".")[0]: p.runs[0].font.size.pt for p in prompts}
            self.assertEqual(sizes["2"], 13)
            self.assertEqual(sizes["3"], 18)
            self.assertEqual(sizes["4"], 13)
            self.assertNotIn("Only this question changes", "\n".join(p.text for p in Document(edited).paragraphs))
            formatter.export_survey_config(source, settings, loaded)
            self.assertEqual(json.loads(settings.read_text())["question_overrides"]["QID3"]["_comment"], "Only this question changes")

    def test_cli_default_name_custom_name_layers_and_no_conversion(self):
        for script in ("qsf_to_paper.py", "qsf_logic_to_paper.py"):
            with self.subTest(script=script), tempfile.TemporaryDirectory() as folder:
                folder = Path(folder)
                for name in ("qsf_to_paper.py", "qsf_logic.py", "qsf_logic_to_paper.py"):
                    shutil.copyfile(ROOT / name, folder / name)
                source = folder / "My survey.qsf"
                source.write_text(json.dumps(fixture()))
                result = subprocess.run([sys.executable, str(folder / script), str(source), "--export-survey-config"], capture_output=True, text=True)
                self.assertEqual(result.returncode, 0, result.stderr)
                default = folder / "My survey_config.json"
                self.assertTrue(default.exists())
                self.assertEqual(json.loads(default.read_text())["question_overrides"]["QID1"]["questions"]["font_size"], 12)
                self.assertFalse(list(folder.glob("*.docx")))
                self.assertFalse(list(folder.glob("*-logic.json")))
                main, survey, output = [folder / name for name in ("main.json", "overrides.json", "custom.json")]
                main.write_text(json.dumps({"questions": {"font_size": 13}}))
                survey.write_text(json.dumps({"questions": {"font_size": 14}, "question_overrides": {"QID2": {"questions": {"font_size": 16}}}}))
                result = subprocess.run([sys.executable, str(folder / script), str(source), "--config", str(main),
                    "--survey-config", str(survey), "--set", "questions.font_size=15", "--export-survey-config", str(output)], capture_output=True, text=True)
                self.assertEqual(result.returncode, 0, result.stderr)
                result_config = json.loads(output.read_text())
                self.assertEqual(result_config["question_overrides"]["QID1"]["questions"]["font_size"], 15)
                self.assertEqual(result_config["question_overrides"]["QID2"]["questions"]["font_size"], 16)

    def test_real_survey_round_trip_and_repeatable_export(self):
        source = ROOT / "examples/Example_survey.qsf"
        config, warnings = formatter.load_config_layers(survey_path=ROOT / "examples/Example_survey_config.json")
        self.assertFalse(warnings)
        with tempfile.TemporaryDirectory() as folder:
            folder = Path(folder)
            settings, original, reloaded = [folder / name for name in ("settings.json", "original.docx", "reloaded.docx")]
            path, count, warnings = formatter.export_survey_config(source, settings, config)
            self.assertEqual(count, 22)
            self.assertFalse(warnings)
            snapshot = json.loads(settings.read_text())
            loaded, warnings = formatter.load_config_layers(survey_path=settings)
            self.assertFalse(warnings)
            for converter in (formatter.convert, adapter.convert_with_logic):
                converter(source, original, config=config)
                converter(source, reloaded, config=loaded)
                self.assertEqual(Document(original).element.xml, Document(reloaded).element.xml)
            formatter.export_survey_config(source, settings, loaded)
            self.assertEqual(json.loads(settings.read_text()), snapshot)

    def test_invalid_sources_conflicting_modes_and_source_protection(self):
        with tempfile.TemporaryDirectory() as folder:
            folder = Path(folder)
            source = folder / "survey.qsf"
            source.write_text(json.dumps(fixture()))
            original = source.read_bytes()
            with self.assertRaisesRegex(ValueError, "overwrite the source"):
                formatter.export_survey_config(source, source)
            self.assertEqual(source.read_bytes(), original)
            for script in ("qsf_to_paper.py", "qsf_logic_to_paper.py"):
                for extra in (["--list-questions"], ["-o", str(folder / "out.docx")]):
                    result = subprocess.run([sys.executable, str(ROOT / script), str(source), "--export-survey-config", *extra], capture_output=True, text=True)
                    self.assertNotEqual(result.returncode, 0)
                    self.assertNotIn("Traceback", result.stderr)
            for data in ({}, {"SurveyElements": []}, [], None):
                source.write_text(json.dumps(data))
                with self.assertRaises(ValueError): formatter.export_survey_config(source)
        formatter.resolve_config({"question_overrides": {"QID1": {"_comment": ""}}})
        with self.assertRaises(ValueError): formatter.resolve_config({"question_overrides": {"QID1": {"_comment": 12}}})
        with self.assertRaises(ValueError): formatter.resolve_config({"_comment": "Not a global setting"})


if __name__ == "__main__":
    unittest.main()
