"""Check response geometry, long-list styling, and handwriting space."""
import json
import sys
import unittest
from pathlib import Path

from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import qsf_to_paper as m


def choices(count):
    return {str(i): {"Display": "Option {}".format(i)} for i in range(1, count + 1)}


def choice_question(count=2):
    return {"QuestionType": "MC", "Selector": "SAVR", "QuestionText": "Choose a response",
            "Choices": choices(count)}


class ResponseLayoutTests(unittest.TestCase):
    def test_yes_no_autoscale_is_vertical_and_ratings_stay_horizontal(self):
        q = choice_question()
        q["Choices"] = {"1": {"Display": "<b>Yes</b>"}, "2": {"Display": "No"}}
        q["Configuration"] = {"Autoscale": {"YScale": {"Name": "yesNo", "Type": "likert"}}}
        w = m.Writer("Test")
        w.question("QID1", q)
        table = w.doc.tables[0]
        self.assertEqual([[c.text for c in r.cells] for r in table.rows], [["◯", "Yes"], ["◯", "No"]])
        q["Choices"] = {str(i): {"Display": label} for i, label in enumerate(
            ["Strongly agree", "Agree", "Neither", "Disagree", "Strongly disagree"], 1)}
        q["Configuration"]["Autoscale"]["YScale"]["Name"] = "agreement"
        w.question("QID2", q)
        self.assertEqual(len(w.doc.tables[1].columns), 5)
        self.assertTrue(all(c.text == "◯" for c in w.doc.tables[1].rows[1].cells))

    def test_indented_tables_fit_printable_width_in_both_layouts(self):
        for layout in ("vertical", "horizontal"):
            for indent in (0, .5, 1):
                w = m.Writer("Test", config={"responses": {"layout": layout, "indent_inches": indent}})
                w.question("QID1", choice_question(5))
                table = w.doc.tables[0]
                node = table._tbl.tblPr.find(qn("w:tblInd"))
                self.assertEqual(int(node.get(qn("w:w"))), round(indent * 1440))
                self.assertAlmostEqual(sum(c.width.inches for c in table.columns) + indent, w.width, places=2)
        with self.assertRaisesRegex(ValueError, "indentation"):
            m.resolve_config({"responses": {"indent_inches": 10}})

    def test_large_list_threshold_banding_and_column_layout(self):
        for count, banded in ((10, False), (11, True), (18, True)):
            w = m.Writer("Test")
            w.question("QID1", choice_question(count))
            table = w.doc.tables[0]
            self.assertEqual(table.cell(1, 1)._tc.find(".//" + qn("w:shd")) is not None, banded)
            self.assertEqual(len(table.columns), 5 if count > 10 else 2)
            labels = [c.text for r in table.rows for c in r.cells if c.text.startswith("Option")]
            self.assertCountEqual(labels, ["Option {}".format(i) for i in range(1, count + 1)])
        q = choice_question(11)
        for choice in q["Choices"].values():
            choice["Display"] += " with a longer explanatory label that stays vertical"
        w = m.Writer("Test")
        w.question("QID1", q)
        self.assertEqual(len(w.doc.tables[0].columns), 2)
        self.assertIsNotNone(w.doc.tables[0].cell(1, 1)._tc.find(".//" + qn("w:shd")))
        w = m.Writer("Test", config={"responses": {"large_list_banded_rows": False}})
        w.question("QID1", q)
        self.assertIsNone(w.doc.tables[0].cell(1, 1)._tc.find(".//" + qn("w:shd")))
        w = m.Writer("Test", config={"responses": {"layout": "horizontal"}})
        w.question("QID1", choice_question(11))
        self.assertIsNotNone(w.doc.tables[0].cell(1, 1)._tc.find(".//" + qn("w:shd")))

    def test_form_rows_expand_only_within_page_budget(self):
        # A form should offer more handwriting space than a compact rank list,
        # then reduce the added space when a whole question would not fit.
        for budget, expected in ((100, .45), (54, .375), (40, .34)):
            w = m.Writer("Test")
            w.fields(["First field", "Second field"], available_height=budget)
            table = w.doc.tables[0]
            self.assertTrue(all(abs(row.height.inches - expected) < .002 for row in table.rows))
        w = m.Writer("Test")
        w.fields(["First rank", "Second rank"], answer_first=True)
        self.assertTrue(all(abs(row.height.inches - .34) < .002 for row in w.doc.tables[0].rows))
        w = m.Writer("Test")
        w.fields(["First field", "Second field"], available_height=100, current_page_height=54)
        self.assertTrue(all(abs(row.height.inches - .375) < .002 for row in w.doc.tables[0].rows))
        w = m.Writer("Test")
        w.fields(["First field", "Second field"], available_height=100, current_page_height=40)
        self.assertTrue(all(abs(row.height.inches - .45) < .002 for row in w.doc.tables[0].rows))
        w = m.Writer("Test")
        w.question("QID1", {"QuestionType": "TE", "Selector": "FORM", "QuestionText": "Comments",
                             "Choices": choices(17)})
        self.assertTrue(all(abs(row.height.inches - .45) < .002 for row in w.doc.tables[0].rows))
        self.assertEqual(w.doc.tables[0]._tbl.tblPr.find(qn("w:tblInd")).get(qn("w:w")), "0")

    def test_form_growth_uses_remaining_space_after_existing_responses(self):
        w = m.Writer("Test")
        w.lines(17)
        w.question("QID1", {"QuestionType": "TE", "Selector": "FORM", "QuestionText": "Comments",
                             "Choices": choices(8)})
        self.assertTrue(all(.34 <= row.height.inches < .45 for row in w.doc.tables[1].rows))

    def test_defaults_and_per_question_overrides(self):
        self.assertEqual(json.loads((ROOT / "config.json").read_text()), m.DEFAULT_CONFIG)
        w = m.Writer("Test", config={"question_overrides": {"QID1": {
            "responses": {"indent_inches": 0, "large_list_banded_rows": False,
                          "form_preferred_row_height_inches": .6}}}})
        w.question("QID1", choice_question(11))
        w.question("QID2", choice_question(11))
        self.assertEqual(w.doc.tables[0]._tbl.tblPr.find(qn("w:tblInd")).get(qn("w:w")), "0")
        self.assertIsNone(w.doc.tables[0].cell(1, 1)._tc.find(".//" + qn("w:shd")))
        self.assertEqual(w.doc.tables[1]._tbl.tblPr.find(qn("w:tblInd")).get(qn("w:w")), "720")
        self.assertIsNotNone(w.doc.tables[1].cell(1, 1)._tc.find(".//" + qn("w:shd")))


if __name__ == "__main__":
    unittest.main()
