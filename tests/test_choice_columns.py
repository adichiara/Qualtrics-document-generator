"""Check row caps, manual option columns, and content-preserving fallbacks."""
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import qsf_to_paper as formatter
import qsf_logic_to_paper as adapter
from test_dropdown_layout import question, table_labels


RANKS = (["E-{}".format(i) for i in range(1, 10)] + ["WO-1"] +
         ["CWO-{}".format(i) for i in range(2, 6)] + ["O-{}".format(i) for i in range(1, 11)])


def columns(table):
    return [[table.cell(row, column).text for row in range(len(table.rows))
             if table.cell(row, column).text] for column in range(1, len(table.columns), 3)]


class ChoiceColumnTests(unittest.TestCase):
    def writer(self, settings=None, labels=RANKS, selector="DL"):
        writer = formatter.Writer("Column checks", config={"responses": settings or {}})
        q = question(labels); q["Selector"] = selector
        writer.question("QID1", q)
        return writer

    def test_rank_families_group_by_default_independently_of_qid(self):
        writer = formatter.Writer("Automatic rank grouping")
        for qid in ("QID1", "QID19789337", "QID999999"):
            writer.question(qid, question(RANKS))
        for table in writer.doc.tables:
            self.assertEqual(columns(table), [RANKS[:9], RANKS[9:14], RANKS[14:]])
            self.assertEqual(len(table.rows), 10)
        self.assertFalse(writer.warnings)
        writer = self.writer({"group_rank_grades": False})
        self.assertEqual([len(c) for c in columns(writer.doc.tables[0])], [8, 8, 8])
        writer = self.writer({"column_breaks_before": [6, 15]})
        self.assertEqual([len(c) for c in columns(writer.doc.tables[0])], [5, 9, 10])
        decorated = ["<b>" + rank + "</b>" for rank in RANKS]
        writer = self.writer(labels=decorated)
        self.assertEqual(columns(writer.doc.tables[0]), [RANKS[:9], RANKS[9:14], RANKS[14:]])
        # Family boundaries, rather than hard-coded positions, handle fewer grades.
        shortened = ["E-5", "E-6", "W-1", "W-2", "O-2", "O-3"]
        writer = self.writer(labels=shortened)
        self.assertEqual(columns(writer.doc.tables[0]), [shortened[:2], shortened[2:4], shortened[4:]])
        # Ambiguous/noncontiguous family order keeps ordinary list behavior.
        writer = self.writer(labels=["E-1", "O-1", "E-2"])
        self.assertEqual(len(writer.doc.tables[0].columns), 2)

    def test_rank_groups_by_positions_or_labels(self):
        for breaks in ([10, 15], ["WO-1", "O-1"], [15, 10, 10], ["WO-1", 15]):
            with self.subTest(breaks=breaks):
                writer = self.writer({"max_rows": 10, "column_breaks_before": breaks})
                self.assertFalse(writer.warnings)
                table = writer.doc.tables[0]
                self.assertEqual(len(table.rows), 10)
                self.assertEqual(columns(table), [RANKS[:9], RANKS[9:14], RANKS[14:]])
                self.assertEqual(table.cell(9, 0).text, "")
                self.assertEqual(table.cell(9, 1).text, "")
                for row in range(5, 10):
                    self.assertEqual(table.cell(row, 3).text, "")
                    self.assertEqual(table.cell(row, 4).text, "")
                self.assertEqual(table.cell(9, 6).text, "◯")
                self.assertEqual(table.cell(9, 7).text, "O-10")
                self.assertAlmostEqual(sum(c.width.inches for c in table.columns) + .5, writer.width, places=2)

    def test_max_rows_adds_columns_and_further_splits_explicit_groups(self):
        writer = self.writer({"max_rows": 5})
        self.assertEqual(len(writer.doc.tables[0].rows), 5)
        self.assertEqual(table_labels(writer.doc.tables[0]), RANKS)
        writer = self.writer({"max_rows": 7, "column_breaks_before": [10, 15]})
        self.assertEqual([len(c) for c in columns(writer.doc.tables[0])], [7, 2, 5, 7, 3])
        self.assertEqual(table_labels(writer.doc.tables[0]), RANKS)
        # A row cap can split even a short or long-label list when explicitly set.
        writer = self.writer({"max_rows": 2}, labels=["A longer descriptive choice {}".format(i) for i in range(6)])
        self.assertEqual(len(writer.doc.tables[0].rows), 2)

    def test_panels_preserve_breaks_rows_width_and_font_size(self):
        writer = self.writer({"max_rows": 3})
        self.assertEqual(len(writer.doc.tables), 2)
        self.assertEqual([label for table in writer.doc.tables for label in table_labels(table)], RANKS)
        sizes = set()
        for table in writer.doc.tables:
            self.assertLessEqual(len(table.rows), 3)
            self.assertLessEqual(sum(c.width.inches for c in table.columns) + .5, writer.width + .01)
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        self.assertEqual(paragraph.paragraph_format.space_before.pt, 0)
                        self.assertEqual(paragraph.paragraph_format.space_after.pt, 0)
                        for run in paragraph.runs:
                            if run.text and run.text != "◯": sizes.add(run.font.size.pt)
        self.assertEqual(len(sizes), 1)
        self.assertTrue(any("QID1:" in message and "stacked table panels" in message for message in writer.warnings))
        # Manual breaks beyond the page capacity are retained in another panel.
        writer = self.writer({"column_breaks_before": [2, 3, 4, 5, 6]}, labels=RANKS[:8])
        self.assertEqual(len(writer.doc.tables), 2)
        self.assertEqual([len(c) for table in writer.doc.tables for c in columns(table)], [1, 1, 1, 1, 1, 3])

    def test_invalid_targets_warn_without_losing_choices(self):
        writer = self.writer({"column_breaks_before": [100, "Missing", "E-1", 10]})
        self.assertEqual(table_labels(writer.doc.tables[0]), RANKS)
        self.assertEqual(len(writer.warnings), 3)
        self.assertEqual(columns(writer.doc.tables[0]), [RANKS[:9], RANKS[9:]])
        writer = self.writer({"column_breaks_before": ["Duplicate"]}, labels=["Duplicate", "Duplicate", "Other"])
        self.assertEqual(table_labels(writer.doc.tables[0]), ["Duplicate", "Duplicate", "Other"])
        self.assertIn("ambiguous", writer.warnings[0])

    def test_independent_per_question_overrides_and_default_layout(self):
        writer = formatter.Writer("Override checks", config={"question_overrides": {"QID1": {
            "responses": {"max_rows": 10, "column_breaks_before": [6, 15]}}}})
        writer.question("QID1", question(RANKS))
        writer.question("QID2", question(RANKS))
        self.assertEqual([len(c) for c in columns(writer.doc.tables[0])], [5, 9, 10])
        self.assertEqual([len(c) for c in columns(writer.doc.tables[1])], [9, 5, 10])
        # Explicit [] clears inherited manual breaks rather than appending them.
        writer = formatter.Writer("Override checks", config={"responses": {"column_breaks_before": [6, 15]},
            "question_overrides": {"QID1": {"responses": {"column_breaks_before": []}}}})
        writer.question("QID1", question(RANKS))
        self.assertEqual([len(c) for c in columns(writer.doc.tables[0])], [9, 5, 10])

    def test_ratings_and_multiple_selection_keep_response_semantics(self):
        q = question(["Strongly agree", "Agree", "Neutral", "Disagree", "Strongly disagree"])
        writer = formatter.Writer("Ratings", config={"responses": {"max_rows": 1}})
        writer.question("QID1", q)
        self.assertEqual(len(writer.doc.tables[0].columns), 5)
        self.assertEqual(len(writer.doc.tables[0].rows), 2)
        writer = formatter.Writer("Ratings", config={"responses": {"column_breaks_before": [3]}})
        writer.question("QID1", q)
        self.assertEqual([len(c) for c in columns(writer.doc.tables[0])], [2, 3])
        writer = self.writer({"column_breaks_before": [10, 15]}, selector="MAVR")
        self.assertEqual(writer.doc.tables[0].cell(0, 0).text, "☐")

    def test_validation_layers_and_cli(self):
        self.assertEqual(json.loads((ROOT / "config.json").read_text()), formatter.DEFAULT_CONFIG)
        for settings in ({"max_rows": 0}, {"max_rows": 2.5}, {"max_rows": True},
                         {"column_breaks_before": None}, {"column_breaks_before": "WO-1"},
                         {"column_breaks_before": [1]}, {"column_breaks_before": [True]},
                         {"column_breaks_before": [2.5]}, {"column_breaks_before": [""]}):
            with self.assertRaises(ValueError): formatter.resolve_config({"responses": settings})
        with tempfile.TemporaryDirectory() as folder:
            folder = Path(folder)
            qsf, out, main, survey = [folder / name for name in ("survey.qsf", "out.docx", "main.json", "survey.json")]
            data = {"SurveyEntry": {"SurveyName": "Rank columns"}, "SurveyElements": [
                {"Element": "SQ", "Payload": dict(question(RANKS), QuestionID="QID1")}]}
            qsf.write_text(json.dumps(data))
            main.write_text(json.dumps({"responses": {"max_rows": 12, "column_breaks_before": [5]}}))
            survey.write_text(json.dumps({"question_overrides": {"QID1": {
                "responses": {"max_rows": 10, "column_breaks_before": [10, 15]}}}}))
            result = subprocess.run([sys.executable, str(ROOT / "qsf_logic_to_paper.py"), str(qsf), "-o", str(out),
                                     "--config", str(main), "--survey-config", str(survey)], capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertEqual([len(c) for c in columns(Document(out).tables[0])], [9, 5, 10])
            result = subprocess.run([sys.executable, str(ROOT / "qsf_to_paper.py"), str(qsf), "-o", str(out),
                                     "--set", "responses.max_rows=10", "--set", 'responses.column_breaks_before=["WO-1","O-1"]'], capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertEqual([len(c) for c in columns(Document(out).tables[0])], [9, 5, 10])


if __name__ == "__main__":
    unittest.main()
