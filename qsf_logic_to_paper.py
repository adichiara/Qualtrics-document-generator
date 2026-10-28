"""Optional Word adapter for qsf_logic; leaves the original converter intact.

python qsf_logic_to_paper.py survey.qsf -o survey-paper-logic.docx
"""
import argparse
import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.table import Table
from docx.text.paragraph import Paragraph

import qsf_logic as logic
import qsf_to_paper as formatter


def add_instruction(doc, anchor, wording, settings, before=False):
    node = OxmlElement("w:p")
    (anchor.addprevious if before else anchor.addnext)(node)
    p = Paragraph(node, doc)
    p.paragraph_format.space_before = Pt(0 if before else 3)
    p.paragraph_format.space_after = Pt(3 if before else 0)
    p.paragraph_format.line_spacing = 1
    p.paragraph_format.keep_together = True
    p.paragraph_format.keep_with_next = before
    run = p.add_run(wording)
    run.bold = True
    run.font.name = settings["fonts"]["family"]
    run.font.size = Pt(settings["responses"]["font_size"])
    return p


def convert_with_logic(qsf, output, config=None, report_path=None):
    config = formatter.resolve_config(config)
    with open(qsf, encoding="utf-8-sig") as source:
        data = json.load(source)
    questions = {e["Payload"]["QuestionID"]: e["Payload"] for e in data.get("SurveyElements", [])
                 if e.get("Element") == "SQ" and isinstance(e.get("Payload"), dict)}
    modes = {qid: config["question_overrides"].get(qid, {}).get("questions", {}).get(
                 "numbering", config["questions"]["numbering"]) for qid in questions}
    report = logic.analyze_survey(data, config["questions"]["numbering"], modes)
    if not config["sections"]["show_block_names"]:
        # A section name cannot serve as a paper skip destination when hidden.
        blocks = report["blocks"]
        for instruction in report["instructions"]:
            if instruction["kind"] not in {"skip", "conditional_skip"} or instruction.get("target_question_id"):
                continue
            current = next((i for i, block in enumerate(blocks)
                            if block["block_id"] == instruction["block_id"]), None)
            if current is None:
                continue
            target = next((qid for block in blocks[current + 1:] for qid in block["question_ids"]
                           if qid in report["question_references"]), None)
            direction = "Skip to {}.".format(report["question_references"][target]) if target else "Skip the remaining questions."
            instruction["text"] = re.sub(
                r"(?:Skip|skip) (?:to the .* section|the remaining questions in this section)\.$",
                lambda match: direction if match[0][0].isupper() else direction[0].lower() + direction[1:],
                instruction["text"])
            instruction["target_question_id"] = target
    count, warnings = formatter.convert(qsf, output, config=config)
    if report["instructions"] and any(mode == "off" for mode in modes.values()):
        warnings.append("Some question numbers are hidden; ensure all routing references are identifiable before printing.")
    doc = Document(output)
    ordered_ids = [qid for block in report["blocks"] for qid in block["question_ids"]]
    anchors, cursor, number = {}, 0, 0
    # Locate only body prompts in survey order; response-table labels cannot match.
    for qid in ordered_ids:
        q = questions[qid]
        is_description = q.get("QuestionType") == "DB"
        if not is_description:
            number += 1
        mode = modes[qid]
        tag = formatter.clean(q.get("DataExportTag")) if mode == "survey" else str(number)
        tag = tag or str(number)
        first_line = str(formatter.clean(q.get("QuestionText"))).split("\n")[0]
        prefix = ("{}.  ".format(tag) if mode != "off" and not is_description else "") + first_line
        if not prefix:
            continue
        for position in range(cursor, len(doc.paragraphs)):
            p = doc.paragraphs[position]
            if p.text.startswith(prefix):
                anchors[qid] = p
                cursor = position + 1
                break
        else:
            raise ValueError("Could not locate {} in the generated document.".format(qid))
    all_prompt_nodes = {p._p for p in anchors.values()}
    instruction_ends = {}

    def settings_for(qid):
        return formatter.merge_config(config, config["question_overrides"].get(qid, {}))

    def response_elements(qid):
        nodes = []
        node = anchors[qid]._p.getnext()
        while node is not None and node not in all_prompt_nodes:
            if node.tag == qn("w:tbl"):
                nodes.append(node)
            elif node.tag == qn("w:p"):
                p = Paragraph(node, doc)
                if not p.text.strip() or p.style.name.startswith("Heading"):
                    break
                nodes.append(node)
            else:
                break
            node = node.getnext()
        return nodes

    for instruction in report["instructions"]:
        kind = instruction["kind"]
        if kind == "block_eligibility":
            block = next(b for b in report["blocks"] if b["block_id"] == instruction["block_id"])
            qid = block["question_ids"][0]
        else:
            qid = instruction["question_id"]
        if qid not in anchors:
            raise ValueError("No insertion point for {}.".format(qid))
        anchor, settings = anchors[qid], settings_for(qid)
        if kind in {"question_eligibility", "block_eligibility"}:
            add_instruction(doc, anchor._p, instruction["text"], settings, before=True)
            continue
        nodes = response_elements(qid)
        if kind == "skip":
            label = logic.text(formatter.clean(questions[qid].get("Choices", {}).get(instruction["choice_id"])))
            matching = []
            for node in nodes:
                if node.tag == qn("w:tbl"):
                    for row in Table(node, doc).rows:
                        for cell in row.cells:
                            if logic.text(cell.text) == label:
                                matching.append(cell)
            # Attach to the label cell, keeping mark and response label separate.
            if len(matching) == 1:
                cell = matching[0]
                p = cell.paragraphs[0]
                size = next((r.font.size for r in p.runs if r.font.size is not None), Pt(settings["responses"]["font_size"]))
                run = p.add_run(" — " + instruction["text"])
                run.bold = True
                run.font.name = settings["fonts"]["family"]
                run.font.size = size
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                continue
        # Dropdowns converted to write-ins, or ambiguous label cells, receive
        # question-level directions after the response area.
        wording = instruction["text"]
        if kind == "skip":
            wording = 'If you selected “{}”, {}{}'.format(instruction["choice_label"], wording[0].lower(), wording[1:])
        end = instruction_ends.get(qid, nodes[-1] if nodes else anchor._p)
        if settings["pagination"]["keep_questions_together"]:
            if end.tag == qn("w:tbl"):
                for cell in Table(end, doc).rows[-1].cells:
                    for p in cell.paragraphs:
                        p.paragraph_format.keep_with_next = True
            else:
                Paragraph(end, doc).paragraph_format.keep_with_next = True
        p = add_instruction(doc, end, wording, settings)
        instruction_ends[qid] = p._p
    doc.save(output)
    if report_path is not None:
        Path(report_path).write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return count, warnings, report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("qsf", type=Path)
    parser.add_argument("-o", "--output", type=Path)
    parser.add_argument("--config", type=Path)
    parser.add_argument("--survey-config", type=Path)
    parser.add_argument("--numbering", choices=["sequential", "survey", "off"])
    parser.add_argument("--set", action="append", default=[], metavar="SECTION.KEY=VALUE")
    parser.add_argument("--report", type=Path, help="Save the analysis report as JSON")
    parser.add_argument("--list-questions", action="store_true", help="List QSF question IDs, export tags, and prompts, then exit")
    args = parser.parse_args()
    try:
        if args.list_questions:
            formatter.list_questions(args.qsf)
            return 0
        config, config_warnings = formatter.load_config_layers(args.config, args.survey_config)
        cli = {}
        if args.numbering:
            formatter.config_assignment(cli, "questions.numbering=" + args.numbering)
        for assignment in args.set:
            formatter.config_assignment(cli, assignment)
        config = formatter.resolve_config(config, cli)
        output = args.output or args.qsf.with_name(args.qsf.stem + "-paper-logic.docx")
        count, warnings, report = convert_with_logic(args.qsf, output, config, args.report)
        print("Wrote {} ({} response questions; {} respondent instructions).".format(output, count, len(report["instructions"])))
        for warning in config_warnings + warnings:
            print("Review: " + warning, file=sys.stderr)
        for item in report["review_required"]:
            print("Review {}: {}".format(item["source_id"], item["reason"]), file=sys.stderr)
        return 0
    except (OSError, ValueError, TypeError) as exc:
        print("Error: {}".format(exc), file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main())
